"""Extract Sugar's measured hyperfine components and pair them with their lines.

WHAT THIS IS FOR
================
A line of Pr III that is broadened by hyperfine structure (hfs - the splitting
of a level into components labeled by F = I + J, where I = 5/2 is the nuclear
spin of 141Pr and J is the electronic angular momentum) is marked in Sugar's
1974 line table with the character ``*r`` or ``*v``: the line is "degraded"
towards the red (lower wavenumber) or towards the violet (higher wavenumber).

Everything the project has done with those flags so far treats them as a
direction and nothing more, and the size of the displacement is *inferred*
from the fit residuals through the kappa model of `Work_on_hfs_plan.md`
section 2.2.  But Sugar tabulated more than the flag.  For many flagged lines
he also printed the wavelengths of the individual hfs components he could
resolve, as extra rows carrying no intensity.  Those rows are measurements of
the splitting, at his own measurement precision, and they are what this module
recovers.

WHERE THEY COME FROM
====================
The checked extraction of Sugar's Table 1,
``Pr_3_Sugar74_Table1_extracted_v3_gemini-3-flash-preview.xlsm``, sheet
``Table 1``: 4333 rows, one per printed line, with the adopted wavenumber in
the column headed ``wn adopted``.  A component row is one whose ``Intens``
column is 0.

THE ROWS THAT LOOK LIKE COMPONENTS AND ARE NOT
==============================================
``Intens = 0`` is overloaded.  Besides Sugar's own component rows, the
curation of the workbook set the intensity of two rows to zero by hand, each
with a note of the form ``int 20 changed to 0``:

* rows 80 and 2389 - lines Sugar printed with an intensity (20 and 5) whose
  classification duplicated that of the neighboring flagged line, merged into
  it by the curation.  Each lies about 0.1-0.3 cm^-1 on the flagged side of
  its parent, which is to say it behaves exactly like a resolved component;
  they are kept, and marked ``origin = curated_duplicate`` so that nothing
  downstream mistakes them for rows Sugar printed without an intensity.

A third, row 2053 (31604.539 cm^-1, a 45.7 cm^-1 Ritz mismatch in Sugar's
table), was zeroed the same way until the curation of 2026-09-23 gave it back
its intensity (4) and character (``c``) as an unclassified line.  Were it ever
zeroed again it would be dropped, not kept: it is 6.9 cm^-1 from the nearest
flagged line and belongs to no pattern.

So 500 rows carry ``Intens = 0``, all 500 are written here, and 498 of those
are Sugar's own.

HOW A COMPONENT IS PAIRED WITH ITS LINE
=======================================
By wavenumber, to the nearest flagged line within 3 cm^-1, and never by a
rounded key.  Two independent checks say the pairing is right:

* the side.  Every component of an ``*r`` line comes out at LOWER wavenumber
  than the tabulated line and every component of a ``*v`` line at higher, in
  500 cases out of 500, without a single exception.  Nothing in the pairing
  rule enforces that, so it is a test of both the extraction and the reading
  of the flags.
* the alternative rule.  Requiring the flagged side first and only then taking
  the nearest candidate gives the same parent for all 500 rows, although 17
  components have a second flagged line within the 3 cm^-1 window.

The parent line is then joined to the pipeline's own line list,
``line_classifications.csv``, by nearest wavenumber within 0.006 cm^-1 - the
same tolerance the rest of the project joins on.  All 296 flagged lines match
one.  Where that line has exactly one accepted classification its two level
identifiers are written; where it has two or three - a blend - no single pair
of levels owns the pattern, so ``low_id`` and ``upp_id`` are left empty and
the competing pairs are listed in ``class_ids``.  Where it is still
unclassified both are empty and ``n_class`` is 0.

WHAT IS DELIBERATELY NOT DONE HERE
==================================
No fitting.  This module produces the measured positions and nothing derived
from them; the per-pattern fit of the two A constants, and the question of
what feature of the pattern Sugar's tabulated wavelength is, are the next part
of Step 2 of the plan and are kept separate so that the extraction can be
checked on its own.

Run as ``python hfs_components.py`` to rewrite ``hfs_components.csv`` and
print the report.  Nothing in the assignment pipeline is touched.
"""
import argparse
import bisect
import csv
import os
import re
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))

WORKBOOK = 'Pr_3_Sugar74_Table1_extracted_v3_gemini-3-flash-preview.xlsm'
SHEET = 'Table 1'
LINE_LIST = 'line_classifications.csv'
OUTPUT = 'hfs_components.csv'

# Columns of the sheet, 1-based, as the header row names them.
COL = {'wl': 1, 'intens': 3, 'char': 4, 'wn': 31, 'notes': 19, 'note2': 21}

FLAGS = ('*r', '*v')

# A component is paired with the nearest flagged line within this window.  The
# widest displacement actually seen is 2.20 cm^-1.
PAIR_WINDOW = 3.0

# The tolerance on which the project joins one line list to another.
LINE_JOIN_WINDOW = 0.006

# The curation's own marker on a row whose intensity it zeroed by hand.
VOIDED = re.compile(r'int\s+\d+\s+changed to 0')

FIELDS = ['wn_line', 'char', 'wn_component', 'dwn', 'n_components',
          'low_id', 'upp_id', 'wl_line', 'wl_component', 'n_class',
          'class_ids', 'origin', 'xl_row_line', 'xl_row_component']


class ExtractionError(Exception):
    """Raised when the workbook does not hold what this module expects."""


def read_table1(path=None, sheet=SHEET):
    """Return the rows of the extracted Table 1 as dictionaries.

    Each carries `xl_row` (the row number in the sheet, so that any finding
    can be looked up in the workbook by hand), `wl` as printed, `intens`,
    `char`, `wn` (the adopted wavenumber) and the two note columns joined.
    """
    import openpyxl                                  # only needed for a read

    path = path or os.path.join(HERE, WORKBOOK)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ExtractionError('%s has no sheet %r' % (path, sheet))
    ws = wb[sheet]
    out = []
    for xl_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True),
                                    start=2):
        def cell(name):
            i = COL[name] - 1
            return values[i] if i < len(values) else None
        wn = cell('wn')
        if wn is None:
            continue
        notes = ' || '.join(str(cell(k)) for k in ('notes', 'note2')
                            if cell(k) is not None)
        out.append({'xl_row': xl_row,
                    'wl': str(cell('wl') or '').strip(),
                    'intens': cell('intens'),
                    'char': (cell('char') or '').strip(),
                    'wn': float(wn),
                    'notes': notes})
    wb.close()
    return out


def split_zero_intensity(rows):
    """Return (components, voided): the `Intens = 0` rows, curation apart.

    A row the curation zeroed by hand says so in its notes; those are returned
    separately so that the caller decides what to do with each, rather than
    having them silently counted as measurements of Sugar's.
    """
    zero = [r for r in rows if r['intens'] == 0]
    voided = [r for r in zero if VOIDED.search(r['notes'])]
    kept = [r for r in zero if not VOIDED.search(r['notes'])]
    return kept, voided


def pair_with_flagged(components, rows, window=PAIR_WINDOW):
    """Pair each component with the nearest flagged line within `window`.

    Returns (paired, orphans).  `paired` is a list of (component, line, dwn)
    with dwn = wn(component) - wn(line); `orphans` are the components with no
    flagged line in the window.
    """
    flagged = sorted(((r['wn'], r) for r in rows if r['char'] in FLAGS),
                     key=lambda p: p[0])
    keys = [wn for wn, _ in flagged]
    paired, orphans = [], []
    for c in components:
        lo = bisect.bisect_left(keys, c['wn'] - window)
        hi = bisect.bisect_right(keys, c['wn'] + window)
        if lo == hi:
            orphans.append(c)
            continue
        wn, line = min(flagged[lo:hi], key=lambda p: abs(p[0] - c['wn']))
        paired.append((c, line, c['wn'] - wn))
    return paired, orphans


def side_violations(paired):
    """Return the pairs whose component is on the wrong side of its line.

    A component of an `*r` line must lie at lower wavenumber and one of a `*v`
    line at higher.  The list is expected to be empty; it is the standing test
    that the flags and the extraction agree.
    """
    return [p for p in paired
            if (p[1]['char'] == '*r' and p[2] >= 0)
            or (p[1]['char'] == '*v' and p[2] <= 0)]


def read_line_list(path=None):
    """Return {wavenumber: [classification rows]} from the pipeline's list."""
    path = path or os.path.join(HERE, LINE_LIST)
    by_wn = {}
    with open(path, encoding='utf-8', newline='') as fh:
        for row in csv.DictReader(fh):
            if not row.get('wn_obs'):
                continue
            by_wn.setdefault(float(row['wn_obs']), []).append(row)
    return by_wn


def classifications_of(wn, keys, by_wn, window=LINE_JOIN_WINDOW):
    """Return (matched wavenumber, accepted classifications) for a line.

    `matched` is None when no line of the pipeline's list lies within the
    window, which would be a broken join and is reported by the caller.  An
    empty list with a matched wavenumber means the line is still unclassified.
    """
    lo = bisect.bisect_left(keys, wn - window)
    hi = bisect.bisect_right(keys, wn + window)
    if lo == hi:
        return None, []
    best = min(keys[lo:hi], key=lambda k: abs(k - wn))
    accepted = [r for r in by_wn[best]
                if r.get('accepted') in ('1', '1.0', 'True', 'TRUE')
                and r.get('low_id')]
    return best, accepted


def build(path=None, line_list=None):
    """Do the whole extraction; return (records, report)."""
    rows = read_table1(path)
    components, voided = split_zero_intensity(rows)
    kept_voided, dropped = pair_with_flagged(voided, rows)
    kept_voided = [c for c, _, _ in kept_voided]

    paired, orphans = pair_with_flagged(components + kept_voided, rows)
    bad_side = side_violations(paired)

    by_wn = read_line_list(line_list)
    keys = sorted(by_wn)
    counts = Counter(line['xl_row'] for _, line, _ in paired)

    records, unmatched, blends = [], [], []
    for component, line, dwn in sorted(paired,
                                       key=lambda p: (p[1]['wn'], p[0]['wn'])):
        matched, accepted = classifications_of(line['wn'], keys, by_wn)
        if matched is None:
            unmatched.append(line)
        if len(accepted) > 1:
            blends.append((line['xl_row'], accepted))
        one = accepted[0] if len(accepted) == 1 else None
        records.append({
            'wn_line': '%.6f' % line['wn'],
            'char': line['char'],
            'wn_component': '%.6f' % component['wn'],
            'dwn': '%+.4f' % dwn,
            'n_components': counts[line['xl_row']],
            'low_id': one['low_id'] if one else '',
            'upp_id': one['upp_id'] if one else '',
            'wl_line': line['wl'],
            'wl_component': component['wl'],
            'n_class': len(accepted),
            'class_ids': (';'.join('%s+%s' % (r['low_id'], r['upp_id'])
                                   for r in accepted)
                          if len(accepted) > 1 else ''),
            'origin': ('curated_duplicate' if component in kept_voided
                       else 'sugar_component'),
            'xl_row_line': line['xl_row'],
            'xl_row_component': component['xl_row'],
        })

    report = {'rows': len(rows),
              'zero_intensity': len(components) + len(voided),
              'sugar_components': len(components),
              'curated_kept': len(kept_voided),
              'curated_dropped': dropped,
              'orphans': orphans,
              'side_violations': bad_side,
              'flagged_lines': sum(1 for r in rows if r['char'] in FLAGS),
              'lines_with_components': len(counts),
              'unmatched_lines': unmatched,
              'blends': blends}
    return records, report


def write_csv(records, path=None):
    """Write `hfs_components.csv`, LF endings, as the repository requires."""
    path = path or os.path.join(HERE, OUTPUT)
    with open(path, 'w', encoding='utf-8', newline='\n') as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator='\n')
        writer.writeheader()
        writer.writerows(records)
    return path


def _report(report, records):
    """Print what was found, including every check that has to be watched."""
    print('Table 1 rows read                        %5d' % report['rows'])
    print('rows with Intens = 0                     %5d'
          % report['zero_intensity'])
    print("   Sugar's own component rows            %5d"
          % report['sugar_components'])
    print('   curation zeroed, kept as components   %5d'
          % report['curated_kept'])
    print('   curation zeroed, dropped              %5d'
          % len(report['curated_dropped']))
    for r in report['curated_dropped']:
        print('      sheet row %d, %s A, %.3f cm^-1: %s'
              % (r['xl_row'], r['wl'], r['wn'], r['notes'][:70]))
    print('components with no flagged line within %.1f cm^-1   %d'
          % (PAIR_WINDOW, len(report['orphans'])))
    print('components on the wrong side of their line          %d'
          % len(report['side_violations']))
    print('flagged (*r/*v) lines in the table       %5d'
          % report['flagged_lines'])
    print('   of them carrying components           %5d'
          % report['lines_with_components'])
    print('parent lines missing from %s: %d'
          % (LINE_LIST, len(report['unmatched_lines'])))
    print('parent lines that are blends (2+ accepted classifications): %d'
          % len({x for x, _ in report['blends']}))
    multiplicity = Counter(int(r['n_components']) for r in records)
    print('components per line: '
          + ', '.join('%d -> %d lines' % (n, c // n)
                      for n, c in sorted(multiplicity.items())))
    for flag in FLAGS:
        d = [float(r['dwn']) for r in records if r['char'] == flag]
        if d:
            print('%s: %3d components, displacement %+.2f to %+.2f, mean %+.3f'
                  % (flag, len(d), min(d), max(d), sum(d) / len(d)))
    print('component rows with no single level pair: %d unclassified, '
          '%d on a blend'
          % (sum(1 for r in records if r['n_class'] == 0),
             sum(1 for r in records if r['n_class'] > 1)))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--workbook', default=None, help='the extracted Table 1')
    ap.add_argument('--line-list', default=None,
                    help='line_classifications.csv')
    ap.add_argument('--out', default=None, help='output csv')
    ap.add_argument('--dry-run', action='store_true',
                    help='report only, write nothing')
    args = ap.parse_args(argv)

    records, report = build(args.workbook, args.line_list)
    _report(report, records)
    if args.dry_run:
        print('\n--dry-run: nothing written')
    else:
        print('\nwrote %s (%d rows)'
              % (write_csv(records, args.out), len(records)))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
