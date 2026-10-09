"""Set 3, the published set: levels and lines at their centers of gravity.

Step 7b of `Work_on_hfs_plan.md`.  Run from LineClass/:

    python make_set3.py                 # writes final/ from iter_hfs/
    python make_set3.py --out DIR       # into another derived set (a test)

WHAT SET 3 IS
=============
iter_hfs/ fits the lines in the head frame: every level stands at its
F = I + J sublevel, E_cg + S, and every line is moved to the strongest
component of its pattern by (1 - kappa) * D (hfs_correction.py).  Set 3 is
the same fit, with the same assignments, expressed at the centers of gravity
(cg): a level at E_cg, a line at the cg of its pattern.  It is what is
published.

It is DERIVED, not classified.  This program writes final/'s table, line
list and report from iter_hfs/'s, and final/lineclass_config.toml says so
(`derived_from`), which stops classify_lines.py, insert_new_level.py,
move_level.py and discard_level.py there (config.require_not_derived).  So
no assignment can differ from the ones reviewed in iter_hfs/, and
classify_lines.py is not touched.  The chain downstream of the table runs on
final/ as on any set: make_LOPT_input.py, lopt.bat LOPT.par, check_sync.py,
sync_IDEN2.py.

THE MOVE OF A LINE (hfs_correction.Model.cg_shift)
==================================================
A line with accepted transitions i, shares BF_i (normalized over the line's
own transitions, as classify_lines.py does; a blended companion's hfs
component is not one of them), was measured at wn_cg + kappa_i * D_i, so

    wn_cg = wn_measured - sum_i BF_i * (kappa_i * D_i + R_i)

with the kappas, the exception registry and the resolved companions exactly
as in iter_hfs/.  R_i takes off in full the S of a level of [hfs]
resolved_levels: the line ends on its F = I + J sublevel.  The identity
with iter_hfs/'s own hfs_shift (hfs_shift - cg_shift = sum BF_i (D_i + R_i))
is checked on every line, and a line that fails it stops the run: the table
and the model disagree.

Decisions taken with the user (2026-10-09):

* every line, flagged or not, is moved by its D from the A table - head - D
  for a flagged line.  Sugar's printed components are not used line by line:
  the A constants are fitted to them already, and the per-line scale of the
  few that disagree beyond 3 sigma is listed in the report instead;
* a line without an accepted transition is not moved; neither is a level
  without a determined A (S = 0), which keeps the allowance it carries;
* a c, cl, w, d, h or ch line is given the character hfs when its pattern's
  |D| is at least one effective line width (level_shifts.INSTR_FWHM_A,
  combined with the Doppler width at its wavelength);
* the F = 2 partner of a line of a resolved J = 1/2 level is found here (at
  +-A(I+J) from it), not in hfs_satellites.txt, and merged into it.

WHAT IS WRITTEN (all in the output set, final/ by default)
==========================================================
Pr3_lines_set3.xlsx     every row of iter/'s corrected line list, with the
                        columns COLUMNS added.  own_cg is the line's
                        wavenumber in Set 3; Icor_set3 its intensity, with
                        its merged parts' light; char_set3 its character.
                        A line merged into another is no line of Set 3: its
                        rows are on the sheet `merged`, where merged_into
                        gives the `own` of the line it went into.
line_classifications    iter_hfs/'s table row for row, every row of a line
  .csv / .xlsx          moved by the line's cg_shift (wn_obs), with
                        hfs_shift = 0 and u_hfs_shift, u_hfs_undet as they
                        were; low_E, upp_E and rwn taken to the centers of
                        gravity, dif_wn_O-C as it was (write_table); and the
                        columns wn_measured, cg_shift, merged_into.  The rows
                        of a merged line are left out.
set3_report.txt         the counts, and every line a rule could not settle
                        or that the user should see.
IDEN2/                  a copy of iter_hfs/IDEN2, made only while the set
                        has none; sync_IDEN2.py brings it to the new fit.
                        Every run takes the merged lines' rows out of its
                        dlv.dat (remove_merged_from_iden2).

After LOPT has been run on the set, `python make_set3.py --levels` writes

levels_set3.csv         every level of the set's fit: E_cg and LOPT's
                        uncertainty, the level's S and u_S, the total, and
                        the source set's E_head - S + S(ground) beside it,
                        which the fit must reproduce.

THE UNCERTAINTIES
=================
An error dS in a level's S (from its A) moves every line of the level the
same way, and the fit takes it up into the level: a line measured at
wn_cg + kappa*D and moved by kappa*D_est stands at
(E_cg + dS)(upper) - (E_cg + dS)(lower) - (1 - kappa)*dD.  Only
(1 - kappa)*u_D is scatter between the lines - the head frame's u_hfs_shift
- and u_S belongs to the level's center of gravity.  So the table keeps the
source set's u_hfs_shift, the fit of Set 3 is the fit of iter_hfs/ moved by
-S level by level (to LOPT's rounding), and levels_set3.csv adds u_S to each
level.  A level's energy is measured from the ground level's center of
gravity, whose u_S (0.021 cm^-1 for A = +0.0305(19)) is common to all of
them and is given once, not added to each.

The line list's u_cg is the other quantity: the uncertainty of one line's
center of gravity on its own, sum BF*(kappa*u_D + u_R) and the kappa term
(hfs_correction.Model.cg_shift), which a published line carries.
"""
import argparse
import csv
import math
import os
import shutil
import sys

import openpyxl
import pandas as pd

import config
import hfs_correction
import hfs_patterns

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_SET = os.path.join(HERE, 'iter_hfs')
OUT_SET = os.path.join(HERE, 'final')
LIST_NAME = 'Pr3_lines_set3.xlsx'
REPORT_NAME = 'set3_report.txt'
LOPT_LEVELS = 'LOPT_output_levels.txt'
GROUND = '059003.000001'
I_SPIN = hfs_correction.I_SPIN

#: the characters that become hfs when the pattern is wide enough
WIDE_CHARS = ('c', 'cl', 'w', 'd', 'h', 'ch')
#: Sugar's marker of a further identification of the line above
REPEAT = '**'
#: the identity hfs_shift - cg_shift = sum BF (D + R), cm^-1
IDENTITY_TOL = 1e-4
#: the measurement precision of a printed component (the component check)
COMPONENT_SIGMA = 0.03
#: the size of a disagreement listed by the component check, in sigma
COMPONENT_LIST_Z = 3.0
#: how far the F = 2 partner of a resolved level's line may lie from where
#: it is predicted, in combined sigma
PAIR_Z = 3.0

#: the sheet of the line list that keeps the lines merged into others
MERGED_SHEET = 'merged'

#: the columns added to the line list
COLUMNS = ('own_cg', 'cg_shift', 'u_cg', 'D_line', 'kappa_line', 'cg_rule',
           'char_set3', 'Icor_set3', 'merged_into')


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------
def key4(wn):
    """The name of a line in this program's dictionaries: its wn_key to four
    decimals, the precision Sugar's own values are compared at elsewhere
    (wavelength_calibration.list_characters)."""
    return round(float(wn), 4)


def base_char(char):
    """Sugar's character without the repeat marker ('w**' -> 'w')."""
    char = (char or '').strip()
    return char[:-len(REPEAT)] if char.endswith(REPEAT) else char


def fwhm_cm(wn):
    """The effective line width at `wn`, cm^-1: the instrumental width and
    the Doppler width combined in wavelength (level_shifts.py's model)."""
    import level_shifts
    lam = 1e8 / wn
    width_A = math.hypot(level_shifts.INSTR_FWHM_A,
                         level_shifts.doppler_fwhm_factor() * lam)
    return width_A * wn * wn / 1e8


def shaded(D):
    """`hfs,l` for a pattern spread to longer wavelengths (D > 0: the head
    lies on the violet side of the cg), `hfs,s` for the other way."""
    return 'hfs,l' if D > 0 else 'hfs,s'


def flag_shade(char):
    """What Sugar's flag says: `*r` red (longer), `*v` violet (shorter)."""
    return {'*r': 'hfs,l', '*v': 'hfs,s'}.get(char)


# ---------------------------------------------------------------------------
# reading
# ---------------------------------------------------------------------------
class Line(object):
    """One observed line of the source table: its rows and what is derived
    from them."""

    def __init__(self, key, rows):
        self.key = key
        self.rows = rows
        first = rows[0]
        self.wn_key = float(first['wn_key'])
        self.wn = float(first['wn_obs'])
        self.unc = float(first['unc_wn_obs'] or 0.0)
        self.char = (first.get('char') or '').strip()
        self.intens = float(first['obs_intens'] or 0.0)
        self.from_companions = float(first.get('obs_intens_hfs') or 0.0)
        self.hfs_shift = float(first.get('hfs_shift') or 0.0)
        acc = [r for r in rows if is_accepted(r)]
        total = sum(float(r['BF'] or 0.0) for r in acc)
        self.components = [
            (r['low_id'].strip(), r['upp_id'].strip(),
             float(r['BF'] or 0.0) / total if total else 1.0 / len(acc))
            for r in acc]
        self.cg = None          # hfs_correction.CgShift
        self.rule = 'none'
        self.char3 = None
        self.intens3 = self.intens
        self.merged_into = None
        self.has_partner = False    # a resolved level's F = 2 line merged in
        self.char_reason = ''       # why char_set3 differs from Sugar's


def is_accepted(row):
    try:
        return float(row.get('accepted') or 0) == 1.0
    except ValueError:
        return False


def read_table(path):
    """`(fieldnames, rows, {key4: Line})` of a classification table, the
    rows as text, in their order."""
    with open(path, encoding='utf-8', newline='') as fh:
        reader = csv.DictReader(fh)
        fields = list(reader.fieldnames)
        rows = list(reader)
    by_line = {}
    for r in rows:
        by_line.setdefault(key4(r['wn_key']), []).append(r)
    return fields, rows, {k: Line(k, v) for k, v in by_line.items()}


def read_list(path, sheet=None):
    """`(header, rows)` of a line list workbook, the rows as lists."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else '' for h in next(it)]
    rows = [list(r) for r in it if any(v is not None for v in r)]
    wb.close()
    return header, rows


def read_lopt_levels(path):
    """`{level_id: energy}` of a LOPT_output_levels.txt."""
    df = pd.read_csv(path, sep='\t', dtype={'Designation': str})
    return dict(zip(df['Designation'].str.strip(), df['Energy'].astype(float)))


# ---------------------------------------------------------------------------
# the rules
# ---------------------------------------------------------------------------
def move_lines(lines, model, report):
    """Give every line its CgShift; check the identity with the table."""
    bad = []
    for ln in lines.values():
        ln.cg = model.cg_shift(ln.char, ln.wn_key, ln.components)
        if ln.components:
            gap = ln.hfs_shift - ln.cg.shift - (ln.cg.D + ln.cg.R)
            if abs(gap) > IDENTITY_TOL:
                bad.append((ln.wn_key, ln.hfs_shift, ln.cg.shift, gap))
            ln.rule = ('resolved' if ln.cg.R else
                       'head-D' if abs(ln.cg.kappa - 1.0) < 1e-12 else
                       'class kappa')
    if bad:
        lines_ = ['make_set3.py: the table\'s hfs_shift disagrees with the '
                  'model on %d line(s) - the table is not the one this '
                  'model writes; rerun the source set\'s chain first:' % len(bad)]
        lines_ += ['  %.4f  hfs_shift %+.4f  cg_shift %+.4f  gap %+.5f' % b
                   for b in bad[:20]]
        raise SystemExit('\n'.join(lines_))
    report['moved'] = sum(1 for ln in lines.values() if ln.components)


def merge_companions(lines, model, report):
    """The resolved companions of hfs_satellites.txt: an unblended one is
    merged into its main line, whose intensity already holds its light; a
    blended one stays a line of its own transitions, less the light it gave
    away; one whose head is not observed stays as it is."""
    given = {}
    by_main = {}
    for sat in model.satellites.rows:
        comp = lines.get(key4(sat.key))
        if comp is None:
            raise SystemExit('make_set3.py: the companion %s of %s is not in '
                             'the table' % (sat.key, sat.main_key))
        if sat.main_key is None:
            comp.char3 = 'hfs'
            report.setdefault('headless', []).append(comp)
            continue
        main = lines.get(key4(sat.main_key))
        if main is None:
            raise SystemExit('make_set3.py: the main line %s of the companion '
                             '%s is not in the table' % (sat.main_key, sat.key))
        by_main.setdefault(main.key, []).append((sat, comp))
    for main_key, group in by_main.items():
        main = lines[main_key]
        whole = [c for s, c in group if not s.blend]
        part = [c for s, c in group if s.blend]
        for c in whole:
            c.merged_into = main.wn_key
            c.intens3 = 0.0
            c.char3 = 'hfs'
            given[c.key] = c.intens
        rest = main.from_companions - sum(c.intens for c in whole)
        if len(part) > 1:
            raise SystemExit('make_set3.py: %s has %d blended companions; the '
                             'light each gave cannot be told apart from the '
                             'table' % (main.wn_key, len(part)))
        for c in part:
            if rest < -1e-6 or rest > c.intens + 1e-6:
                raise SystemExit('make_set3.py: the blended companion %s gave '
                                 '%.3f of its %.3f to %s'
                                 % (c.wn_key, rest, c.intens, main.wn_key))
            c.intens3 = c.intens - max(rest, 0.0)
            given[c.key] = max(rest, 0.0)
            report.setdefault('blended', []).append((c, main, rest))
        report.setdefault('merged', []).extend(
            (c, main, 'companion') for c in whole)
    return given


def pair_resolved(lines, model, listed, report):
    """The F = 2 partners of the lines of resolved J = 1/2 levels.

    A line that ends on the F = I + J sublevel of such a level has a partner
    that ends on the sublevel below, A * (I + J) away: at larger wavenumber
    when the level is the lower one and A > 0.  The partner must be an
    observed line with no accepted transition and no other role here, within
    PAIR_Z combined sigma of the prediction.  One candidate is merged; none,
    or more than one, is listed."""
    taken = {ln.key for ln in lines.values() if ln.merged_into is not None}
    accepted = {ln.key for ln in lines.values() if ln.components}
    for ln in sorted(lines.values(), key=lambda l: l.wn_key):
        if not ln.components or not ln.cg.R:
            continue
        if model.satellites.head_pairs(ln.wn_key):
            report.setdefault('paired_by_registry', []).append(ln)
            continue
        for low, upp, bf in ln.components:
            row = model.exception(ln.wn_key, low, upp)
            unres = row.unresolved if row else ()
            for lid, sign in ((low, 1.0), (upp, -1.0)):
                if lid not in model.resolved or lid in unres:
                    continue
                J, A, u_A, _ = model.levels[lid]
                delta = sign * A * (I_SPIN + J)
                want = ln.wn + delta
                hits = []
                for wn, unc, k in listed:
                    if k in accepted or k in taken or k == ln.key:
                        continue
                    tol = PAIR_Z * math.sqrt(unc ** 2 + ln.unc ** 2
                                             + ((I_SPIN + J) * u_A) ** 2)
                    if abs(wn - want) <= tol:
                        hits.append((abs(wn - want), wn, k))
                rec = (ln, lid, want, sorted(hits))
                if len(hits) == 1:
                    partner = lines.get(hits[0][2])
                    report.setdefault('pairs', []).append(rec)
                    taken.add(hits[0][2])
                    if partner is not None:
                        partner.merged_into = ln.wn_key
                        partner.intens3 = 0.0
                        partner.char3 = 'hfs'
                        ln.intens3 += partner.intens
                        ln.has_partner = True
                    else:
                        report.setdefault('pair_list_only', []).append(hits[0])
                else:
                    report.setdefault('unpaired', []).append(rec)


def characters(lines, model, report):
    """char_set3 of every line not settled by a merge."""
    changed = []
    for ln in lines.values():
        if ln.char3 is not None:
            continue
        D = (ln.cg.D + ln.cg.R) if ln.cg is not None else 0.0
        flag = flag_shade(ln.char)
        if flag:
            ln.char3, ln.char_reason = flag, 'flag'
            if ln.components and D and shaded(D) != flag:
                report.setdefault('flag_sign', []).append((ln, D))
        elif model.satellites.head_pairs(ln.wn_key) and D:
            ln.char3, ln.char_reason = shaded(D), 'companions'
        elif ln.has_partner and D:
            ln.char3, ln.char_reason = shaded(D), 'F = 2 partner'
        elif (ln.char in WIDE_CHARS and ln.components and D
              and abs(D) >= fwhm_cm(ln.wn)):
            ln.char3, ln.char_reason = shaded(D), 'width'
        else:
            ln.char3 = ln.char
        if ln.char3 != ln.char:
            changed.append(ln)
    report['char_changed'] = changed


def unmoved_heads(lines, model, report):
    """Main lines of companions that are not moved: their transition is not
    accepted in the source set."""
    for ln in lines.values():
        if model.satellites.head_pairs(ln.wn_key) and not ln.components:
            report.setdefault('unmoved_heads', []).append(ln)


def component_check(lines, model, components_csv, report):
    """Sugar's printed components against head - D, line by line: the scale
    s of the pattern that his components ask for, and the cg it would give,
    head - s*D, against the one adopted, head - D.  Listed, not used."""
    if not components_csv or not os.path.exists(components_csv):
        return
    comp = pd.read_csv(components_csv)
    out = []
    for wl, g in comp.groupby('wn_line'):
        ln = lines.get(key4(wl))
        if ln is None or not ln.components:
            continue
        low, upp, _ = max(ln.components, key=lambda c: c[2])
        try:
            J1, A1, _, J2, A2, _ = model.pattern_of(low, upp)
        except ValueError:
            continue
        D = hfs_patterns.head_displacement(J1, A1, J2, A2)
        m = g.sort_values('dwn', key=abs)['dwn'].to_numpy()
        n = min(len(m), hfs_patterns.ladder_length(J1, J2))
        if n == 0 or D == 0:
            continue
        p = [hfs_patterns.rung(k + 1, J1, A1, J2, A2) for k in range(n)]
        pp = sum(x * x for x in p)
        if pp == 0:
            continue
        s = sum(a * b for a, b in zip(m[:n], p)) / pp
        u = COMPONENT_SIGMA / math.sqrt(pp) * abs(D)
        d = (s - 1.0) * D
        out.append((ln, n, D, s, d, u))
    report['components'] = out


# ---------------------------------------------------------------------------
# writing
# ---------------------------------------------------------------------------
def level_cg(model, lopt):
    """`{level_id: E_cg}` from head-frame LOPT energies: E - S, the S of a
    resolved level taken whole, and the ground level's cg put at 0."""
    s_ground = model.S(GROUND, pattern=True)[0]
    return {lid: E - model.S(lid, pattern=True)[0] + s_ground
            for lid, E in lopt.items()}


def write_table(fields, rows, lines, model, path_csv, path_xlsx, report):
    """The source table moved to the cg, row for row, without the rows of
    the lines merged into others (`merged_into`): those are no lines of
    Set 3.  A line whose light changed by a merge carries it in obs_intens.

    The level energies are the table's own (the classification's), each
    taken to its center of gravity, E - S + S(ground), and rwn with them.
    dif_wn_O-C is kept as the source wrote it: the line against where the
    component is MEASURED, Ritz + kappa*D + R, which is what
    make_LOPT_input.centroid_unc spreads a blend's components by.  For a
    line of one accepted transition that is exactly wn_obs - rwn here, which
    is checked; in a blend the components' kappa*D differ, and so do the two.
    """
    s_ground = model.S(GROUND, pattern=True)[0]
    out_fields = list(fields) + [c for c in ('wn_measured', 'cg_shift',
                                             'merged_into')
                                 if c not in fields]
    seen = {}
    out = []
    worst = 0.0
    dropped = 0
    for r in rows:
        ln = lines[key4(r['wn_key'])]
        if ln.merged_into is not None:      # a part of another line now
            dropped += 1
            continue
        new = dict(r)
        wn_cg = round(ln.wn + ln.cg.shift, 4)
        other = seen.setdefault(wn_cg, ln.key)
        if other != ln.key:
            raise SystemExit('make_set3.py: the lines %s and %s both come to '
                             '%.4f cm^-1; make_LOPT_input.py would take them '
                             'for one' % (lines[other].wn_key, ln.wn_key,
                                          wn_cg))
        new['wn_measured'] = r['wn_obs']
        new['wn_obs'] = repr(wn_cg)
        new['cg_shift'] = '%.6f' % ln.cg.shift
        new['merged_into'] = ''
        # the line's light with its merged parts (or less the part a blended
        # companion gave away), as the line list has it
        if abs(ln.intens3 - ln.intens) > 1e-9:
            new['obs_intens'] = repr(round(ln.intens3, 6))
        # The line is where LOPT is to see it; what its fit uncertainty
        # owes to hfs is unchanged (see "THE UNCERTAINTIES" above), so
        # u_hfs_shift and u_hfs_undet stay as the source set wrote them.
        if 'hfs_shift' in r:
            new['hfs_shift'] = '0.0'
        low = (r.get('low_id') or '').strip()
        upp = (r.get('upp_id') or '').strip()
        for col, lid in (('low_E', low), ('upp_E', upp)):
            if lid and (r.get(col) or '').strip():
                new[col] = repr(round(float(r[col])
                                      - model.S(lid, pattern=True)[0]
                                      + s_ground, 6))
        if (new.get('low_E') or '').strip() and (new.get('upp_E') or '').strip():
            rwn = float(new['upp_E']) - float(new['low_E'])
            new['rwn'] = repr(round(rwn, 6))
            dif = (r.get('dif_wn_O-C') or '').strip()
            if (len(ln.components) == 1 and is_accepted(r) and dif
                    and dif.lower() != 'nan'):
                worst = max(worst, abs(float(dif) - (wn_cg - rwn)))
        out.append(new)
    with open(path_csv, 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=out_fields, lineterminator='\n')
        w.writeheader()
        w.writerows(out)
    df = pd.read_csv(path_csv, dtype={'low_id': str, 'upp_id': str})
    write_xlsx(df, path_xlsx)
    report['dif_check'] = worst
    report['rows_dropped'] = dropped
    return len(out)


def write_xlsx(df, path):
    """The table as classify_lines.write_output gives it."""
    df.to_excel(path, index=False, engine='openpyxl')
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    fmt = {'wn_obs': '0.0000', 'wn_measured': '0.0000', 'wn_key': '0.0000',
           'unc_wn_obs': '0.000', 'obs_intens': '0.000',
           'obs_intens_hfs': '0.000', 'calc_intens': '0.000',
           'orig_calc_intens': '0.000', 'u_calc': '0.000',
           'intens_from_f': '0.000', 'intens_to_f': '0.000',
           'dif_wn_O-C': '0.000', 'low_E': '0.000', 'upp_E': '0.000',
           'rwn': '0.000', 'BF': '0.000', 'cg_shift': '0.0000'}
    header = {str(c.value): c.column for c in ws[1]}
    for name, f in fmt.items():
        if name in header:
            for row in ws.iter_rows(min_row=2, min_col=header[name],
                                    max_col=header[name]):
                for cell in row:
                    if cell.value is not None:
                        cell.number_format = f
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    wb.close()


def write_list(header, rows, lines, cols, path):
    """iter/'s line list with the Set 3 columns.  A line of the list that no
    table holds is not moved and keeps its character and intensity.  A line
    merged into another is no line of Set 3: its rows go to the sheet
    MERGED_SHEET, which keeps the record of where each went."""
    i_own = header.index(cols['wn_key'])
    i_corr = header.index(cols['wn'])
    i_ch = header.index(cols['character'])
    i_I = header.index(cols['intensity'])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(header + list(COLUMNS))
    gone = wb.create_sheet(MERGED_SHEET)
    gone.append(header + list(COLUMNS))
    for r in rows:
        ln = lines.get(key4(r[i_own])) if r[i_own] is not None else None
        own_corr = float(r[i_corr]) if r[i_corr] is not None else None
        char = (r[i_ch] or '').strip()
        repeat = char.endswith(REPEAT)
        if ln is None or own_corr is None:
            extra = [own_corr, 0.0, 0.0, 0.0, None, 'none',
                     char or None, r[i_I], None]
        else:
            c3 = ln.char3 if ln.char3 is not None else ln.char
            if repeat:
                c3 = (c3 or '') + REPEAT
            extra = [round(own_corr + ln.cg.shift, 4), round(ln.cg.shift, 6),
                     round(ln.cg.u, 6), round(ln.cg.D + ln.cg.R, 6),
                     round(ln.cg.kappa, 6), ln.rule, c3 or None,
                     round(ln.intens3, 4), ln.merged_into]
        merged = ln is not None and ln.merged_into is not None
        (gone if merged else ws).append(list(r) + extra)
    for sheet in (ws, gone):
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    wb.save(path)


def listed_lines(header, rows, cols):
    """`[(wn, unc, key4)]` of the list's lines, the first row of each, on
    the calibrated scale."""
    i_own = header.index(cols['wn_key'])
    i_corr = header.index(cols['wn'])
    i_u = header.index(cols['u_wn'])
    seen, out = set(), []
    for r in rows:
        if r[i_own] is None or r[i_corr] is None:
            continue
        k = key4(r[i_own])
        if k in seen:
            continue
        seen.add(k)
        out.append((float(r[i_corr]), float(r[i_u] or 0.0), k))
    return out


def iden2_line_numbers(trans_path):
    """The dlv.dat line numbers that trans.dat names (its columns 61-66)."""
    out = set()
    with open(trans_path, 'rb') as fh:
        for raw in fh:
            field = raw[60:66].strip()
            if field.isdigit():
                out.add(int(field))
    return out


def remove_merged_from_iden2(iden2, lines, report):
    """Take the rows of the merged lines out of the set's IDEN2: dlv.dat,
    and dlv_keys.txt and dlv_shown.txt, which sync_IDEN2.py keeps beside it.

    A row is found through dlv_keys.txt, by its line's wn_key.  Its line
    number goes with it and is not reused: IDEN2 leaves such gaps itself when
    a line is deleted on its screen, and a number, once given, is how
    trans.dat names the line.  None of these lines may be identified there;
    one that is stops the run before anything is written.  A second run finds
    nothing left to remove."""
    dlv = os.path.join(iden2, 'dlv.dat')
    keys_path = os.path.join(iden2, 'dlv_keys.txt')
    shown = os.path.join(iden2, 'dlv_shown.txt')
    merged = {ln.key: ln for ln in lines.values() if ln.merged_into is not None}
    numbers = {}
    with open(keys_path, encoding='utf-8', newline='') as fh:
        for rec in csv.DictReader(fh, delimiter='\t'):
            k = key4(rec['wn_key'])
            if k in merged:
                numbers[int(rec['line'])] = merged[k]
    named = sorted(n for n in numbers
                   if n in iden2_line_numbers(os.path.join(iden2, 'trans.dat')))
    if named:
        raise SystemExit('make_set3.py: trans.dat identifies the merged '
                         'line(s) %s (dlv.dat line %s); undo that in IDEN2 '
                         'first' % (', '.join('%.4f' % numbers[n].wn_key
                                              for n in named),
                                    ', '.join(map(str, named))))
    report['iden2_removed'] = sorted((n, ln.wn_key)
                                     for n, ln in numbers.items())
    if not numbers:
        return
    import output_files
    output_files.require_writable([dlv, keys_path, shown])

    def last_field(raw):
        parts = raw.split()
        return int(parts[-1]) if parts and parts[-1].isdigit() else None

    def first_field(raw):
        head = raw.split(b'\t', 1)[0].strip()
        return int(head) if head.isdigit() else None

    for path, number_of in ((dlv, last_field), (keys_path, first_field),
                            (shown, first_field)):
        if not os.path.exists(path):
            continue
        with open(path, 'rb') as fh:
            kept = [raw for raw in fh if number_of(raw) not in numbers]
        with open(path, 'wb') as fh:
            fh.writelines(kept)


def write_report(path, report, lines, n_rows, args):
    def f(ln):
        return '%.4f %-4s' % (ln.wn_key, ln.char or '')

    out = ['Set 3 (make_set3.py), from %s' % args.source, '']
    n_lines = len(lines)
    rules = {}
    for ln in lines.values():
        rules[ln.rule] = rules.get(ln.rule, 0) + 1
    out.append('%d lines, %d table rows; moved %d' % (n_lines, n_rows,
                                                     report['moved']))
    out.append('rows of merged lines left out of the table: %d'
               % report.get('rows_dropped', 0))
    out.append('single accepted lines: |dif_wn_O-C - (wn_obs - rwn)| at most '
               '%.1e cm^-1' % report.get('dif_check', 0.0))
    out += ['  rule %-12s %5d' % (k, v) for k, v in sorted(rules.items())]
    shifts = sorted(abs(ln.cg.shift) for ln in lines.values() if ln.components)
    if shifts:
        out.append('  |cg_shift| median %.4f, 90%% %.4f, max %.4f cm^-1'
                   % (shifts[len(shifts) // 2],
                      shifts[int(0.9 * (len(shifts) - 1))], shifts[-1]))
    out.append('')

    merged = report.get('merged', [])
    out.append('Merged companions (hfs_satellites.txt): %d' % len(merged))
    for c, main, _ in merged:
        out.append('  %s -> %s' % (f(c), f(main)))
    out.append('Blended companions, kept as lines: %d'
               % len(report.get('blended', [])))
    for c, main, rest in report.get('blended', []):
        out.append('  %s gave %.1f of %.1f to %s' % (f(c), rest, c.intens,
                                                     f(main)))
    out.append('Companions whose head is not observed (kept, char hfs): %d'
               % len(report.get('headless', [])))
    for c in report.get('headless', []):
        out.append('  %s' % f(c))
    out.append('')

    out.append('Resolved levels: F = 2 partners merged: %d'
               % len(report.get('pairs', [])))
    for ln, lid, want, hits in report.get('pairs', []):
        out.append('  %s (%s) partner %.4f, predicted %.4f'
                   % (f(ln), lid, hits[0][1], want))
    out.append('  lines whose partner is in hfs_satellites.txt: %d'
               % len(report.get('paired_by_registry', [])))
    out.append('  no partner, or more than one: %d'
               % len(report.get('unpaired', [])))
    for ln, lid, want, hits in report.get('unpaired', []):
        out.append('  %s (%s) predicted %.4f: %s'
                   % (f(ln), lid, want, ', '.join('%.4f' % h[1] for h in hits)
                      or 'none'))
    out.append('')

    heads = report.get('unmoved_heads', [])
    out.append('Main lines of companions not moved (transition not accepted '
               'in the source set): %d' % len(heads))
    for ln in heads:
        out.append('  %s' % f(ln))
    out.append('')

    sign = report.get('flag_sign', [])
    out.append('Flags whose side disagrees with the sign of D: %d' % len(sign))
    for ln, D in sign:
        out.append('  %s D %+.4f' % (f(ln), D))
    out.append('')

    ch = report.get('char_changed', [])
    counts = {}
    for ln in ch:
        k = (ln.char, ln.char3, ln.char_reason)
        counts[k] = counts.get(k, 0) + 1
    out.append('Characters changed: %d' % len(ch))
    out += ['  %-4s -> %-6s %5d  (%s)' % (a or "''", b, n, why)
            for (a, b, why), n in sorted(counts.items())]
    out.append('  the wide characters made hfs (|D| >= one line width):')
    for ln in sorted(ch, key=lambda l: l.wn_key):
        if ln.char_reason == 'width':
            D = ln.cg.D + ln.cg.R
            out.append('    %s -> %-6s D %+.4f  width %.4f'
                       % (f(ln), ln.char3, D, fwhm_cm(ln.wn)))
    out.append('')

    comp = report.get('components', [])
    if comp:
        diffs = sorted(abs(c[4]) for c in comp)
        far = [c for c in comp if c[5] > 0 and abs(c[4] / c[5]) > COMPONENT_LIST_Z]
        out.append('Printed components against head - D: %d lines; '
                   '|cg difference| median %.4f, 90%% %.4f cm^-1; beyond '
                   '%.0f sigma: %d' % (len(comp), diffs[len(diffs) // 2],
                                       diffs[int(0.9 * (len(diffs) - 1))],
                                       COMPONENT_LIST_Z, len(far)))
        out.append('  line            n      D       s   head-sD - (head-D)   z')
        for ln, n, D, s, d, u in sorted(far, key=lambda c: -abs(c[4] / c[5])):
            out.append('  %s %2d %+7.4f %7.3f %+9.4f %8.1f'
                       % (f(ln), n, D, s, -d, d / u))
    with open(path, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('\n'.join(out) + '\n')
    return out


# ---------------------------------------------------------------------------
def build(source, out_dir, components_csv=None, copy_iden2=True):
    """Write the derived set `out_dir` from the set `source`; returns the
    report's lines."""
    out_cfg = os.path.join(out_dir, config.CONFIG_NAME)
    if not os.path.isfile(out_cfg) or not config.derived_from(out_dir):
        raise SystemExit('make_set3.py: %s must exist and say derived_from = '
                         '"make_set3.py"' % out_cfg)
    src_cfg = config.load(os.path.join(source, config.CONFIG_NAME))
    if not src_cfg.hfs.apply:
        raise SystemExit('make_set3.py: %s has the hfs correction off; Set 3 '
                         'is derived from a head-frame set' % source)
    dst_cfg = config.load(out_cfg)
    model = hfs_correction.Model(src_cfg.hfs)
    fields, rows, lines = read_table(src_cfg.output_csv)
    header, list_rows = read_list(src_cfg.lines_file, src_cfg.lines.sheet)
    cols = src_cfg.lines.columns

    report = {}
    move_lines(lines, model, report)
    merge_companions(lines, model, report)
    pair_resolved(lines, model, listed_lines(header, list_rows, cols), report)
    characters(lines, model, report)
    unmoved_heads(lines, model, report)
    component_check(lines, model,
                    components_csv if components_csv is not None
                    else getattr(src_cfg.hfs, 'components', ''), report)

    out_list = dst_cfg.lines_file
    for p in (out_list, dst_cfg.output_csv, dst_cfg.output_file):
        if config.set_of(p) is None or os.path.normcase(config.set_of(p)) \
                != os.path.normcase(os.path.abspath(out_dir)):
            raise SystemExit('make_set3.py: %s is not in %s' % (p, out_dir))
    import output_files
    output_files.require_writable([out_list, dst_cfg.output_csv,
                                   dst_cfg.output_file])
    write_list(header, list_rows, lines, cols, out_list)
    n_rows = write_table(fields, rows, lines, model, dst_cfg.output_csv,
                         dst_cfg.output_file, report)
    text = write_report(os.path.join(out_dir, REPORT_NAME), report, lines,
                        n_rows, argparse.Namespace(source=source))
    iden2 = os.path.join(out_dir, 'IDEN2')
    if copy_iden2:
        if not os.path.isdir(iden2) or not os.listdir(iden2):
            # its files, without the backups sync_IDEN2.py left beside them
            shutil.copytree(os.path.join(source, 'IDEN2'), iden2,
                            dirs_exist_ok=True,
                            ignore=shutil.ignore_patterns('*.presync',
                                                          '*.old'))
            text.append('IDEN2 copied from %s'
                        % os.path.join(source, 'IDEN2'))
        remove_merged_from_iden2(iden2, lines, report)
        gone = report['iden2_removed']
        text.append('IDEN2: %d merged line(s) removed from dlv.dat%s'
                    % (len(gone), (' (line ' + ', '.join(
                        '%d %.4f' % g for g in gone) + ')') if gone else ''))
        with open(os.path.join(out_dir, REPORT_NAME), 'a', encoding='utf-8',
                  newline='\n') as fh:
            fh.write('\n' + text[-1] + '\n')
    return text


def write_levels(out_dir, source):
    """levels_set3.csv from the set's own LOPT fit; returns the report's
    lines.  Each level's cg energy and LOPT's uncertainty (the larger of D1
    and D2tot, as sync_IDEN2.py takes it), its S and u_S, the two combined,
    and the source set's energy taken to its cg, with the difference."""
    src_cfg = config.load(os.path.join(source, config.CONFIG_NAME))
    model = hfs_correction.Model(src_cfg.hfs)
    fit = pd.read_csv(os.path.join(out_dir, LOPT_LEVELS), sep='\t',
                      dtype={'Designation': str})
    src = level_cg(model, read_lopt_levels(os.path.join(source, LOPT_LEVELS)))
    rows, diffs = [], []
    for r in fit.itertuples():
        lid = r.Designation.strip()
        u_fit = max(float(r.D1), float(r.D2tot))
        S, u_S = model.S(lid, pattern=True)
        E_src = src.get(lid)
        d = None if E_src is None else float(r.Energy) - E_src
        if d is not None:
            diffs.append((abs(d), lid, d, u_fit))
        rows.append({'level_id': lid, 'E_cg': '%.4f' % float(r.Energy),
                     'u_fit': '%.4f' % u_fit, 'S': '%.4f' % S,
                     'u_S': '%.4f' % u_S,
                     'u_E': '%.4f' % math.hypot(u_fit, u_S),
                     'E_source_cg': '' if E_src is None else '%.4f' % E_src,
                     'diff': '' if d is None else '%.4f' % d,
                     'N_lines': str(r.N_lines).split(',')[0].strip()})
    path = os.path.join(out_dir, 'levels_set3.csv')
    import output_files
    output_files.require_writable([path])
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator='\n')
        w.writeheader()
        w.writerows(rows)
    diffs.sort(reverse=True)
    s_g, u_g = model.S(GROUND, pattern=True)
    out = ['%s: %d levels' % (path, len(rows)),
           'the ground level\'s S = %.4f +- %.4f cm^-1: common to every '
           'energy, not in u_E' % (s_g, u_g),
           'fit - source moved to the cg: largest |diff| %.4f cm^-1 over %d '
           'levels' % (diffs[0][0] if diffs else 0.0, len(diffs))]
    out += ['  %s %+.4f (u_fit %.4f)' % (lid, d, u) for _, lid, d, u
            in diffs[:10]]
    return out


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    p.add_argument('--source', default=SOURCE_SET,
                   help='the head-frame set Set 3 is derived from')
    p.add_argument('--out', default=OUT_SET,
                   help='the derived set written (its lineclass_config.toml '
                        'must say derived_from)')
    p.add_argument('--no-iden2', action='store_true',
                   help='do not copy the source IDEN2 into an empty one')
    p.add_argument('--levels', action='store_true',
                   help='after LOPT has been run on the set: write '
                        'levels_set3.csv from its fit, and nothing else')
    args = p.parse_args(argv)
    if args.levels:
        print('\n'.join(write_levels(os.path.abspath(args.out),
                                     os.path.abspath(args.source))))
        return 0
    text = build(os.path.abspath(args.source), os.path.abspath(args.out),
                 copy_iden2=not args.no_iden2)
    print('\n'.join(text[:12]))
    print('... the full report is in %s'
          % os.path.join(os.path.abspath(args.out), REPORT_NAME))


if __name__ == '__main__':
    sys.exit(main())
