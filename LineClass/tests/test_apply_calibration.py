"""Tests for `tools/apply_calibration.py`.

Run from the LineClass directory:  python -m pytest tests -q

The script has one job: to take the intensity column of a line list from one
wavelength-dependent scale correction to another, without inventing or losing
anything.  A tiny workbook is built here for it to work on.
"""
import os
import sys

import numpy as np
import openpyxl
import pytest

TOOLS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     'tools')
sys.path.insert(0, TOOLS)

import calibrate_intensities as cal      # noqa: E402
import apply_calibration as app          # noqa: E402

SCALE = 1000.0            # the overall factor the line list carries


def make_case(tmp_path, n=200):
    """A workbook whose intensities carry a known correction, plus two files.

    Returns (workbook path, old-correction path, new-correction path, the
    original whole-number intensities, the wavenumbers).
    """
    old = [cal.Region(1000.0, 5000.0, np.polynomial.Polynomial([1.0, -3e-4]))]
    new = [cal.Region(1000.0, 3000.0, np.polynomial.Polynomial([0.5, 1e-4])),
           cal.Region(3000.0, 5000.0, np.polynomial.Polynomial([-2.0, 5e-4]))]
    p_old = str(tmp_path / 'old.txt')
    p_new = str(tmp_path / 'new.txt')
    cal.write_correction(p_old, old)
    cal.write_correction(p_new, new)

    lam = np.linspace(1000.0, 5000.0, n)
    wn = 1e8 / lam
    original = np.round(np.linspace(1.0, 900.0, n))
    stored = SCALE * original * np.exp(cal.apply_correction(old, lam))

    path = str(tmp_path / 'lines.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['own', 'Icor', 'note'])
    for k in range(n):
        ws.append([float(wn[k]), float(stored[k]), 'x'])
    wb.save(path)
    return path, p_old, p_new, original, wn


def read_column(path, name='Icor'):
    ws = openpyxl.load_workbook(path)['Sheet1']
    header = [c.value for c in ws[1]]
    j = header.index(name)
    return np.array([r[j] for r in ws.iter_rows(min_row=2, values_only=True)],
                    float)


def test_the_new_correction_is_the_one_that_ends_up_in_the_column(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    app.main(['--lines', path, '--old', p_old, '--new', p_new,
              '--out-csv', str(tmp_path / 'rec.csv')])
    lam = 1e8 / wn
    want = SCALE * original * np.exp(
        cal.apply_correction(cal.read_correction(p_new), lam))
    assert np.allclose(read_column(path), want, rtol=1e-9)


def test_going_there_and_back_leaves_the_column_as_it_was(tmp_path):
    """old -> new -> old must return the intensities unchanged.

    This is what makes the operation safe to repeat: it is a change of scale,
    not an accumulation.
    """
    path, p_old, p_new, original, wn = make_case(tmp_path)
    before = read_column(path)
    app.main(['--lines', path, '--old', p_old, '--new', p_new,
              '--out-csv', str(tmp_path / 'a.csv')])
    assert not np.allclose(read_column(path), before)
    app.main(['--lines', path, '--old', p_new, '--new', p_old,
              '--out-csv', str(tmp_path / 'b.csv')])
    assert np.allclose(read_column(path), before, rtol=1e-9)


def test_a_dry_run_changes_nothing(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    before = read_column(path)
    app.main(['--lines', path, '--old', p_old, '--new', p_new, '--dry-run',
              '--out-csv', str(tmp_path / 'c.csv')])
    assert np.allclose(read_column(path), before)


def test_the_backup_holds_the_untouched_workbook(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    before = read_column(path)
    app.main(['--lines', path, '--old', p_old, '--new', p_new,
              '--out-csv', str(tmp_path / 'd.csv')])
    bak = os.path.splitext(path)[0] + '.bak.xlsx'
    assert os.path.exists(bak)
    assert np.allclose(read_column(bak), before)


def test_other_columns_and_rows_survive(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    app.main(['--lines', path, '--old', p_old, '--new', p_new,
              '--out-csv', str(tmp_path / 'e.csv')])
    ws = openpyxl.load_workbook(path)['Sheet1']
    assert ws.max_row == len(original) + 1
    assert [c.value for c in ws[1]] == ['own', 'Icor', 'note']
    assert all(r[2] == 'x' for r in ws.iter_rows(min_row=2, values_only=True))


def test_a_missing_column_is_an_error_not_a_silent_pass(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    with pytest.raises(SystemExit):
        app.main(['--lines', path, '--old', p_old, '--new', p_new,
                  '--col-intensity', 'nosuch',
                  '--out-csv', str(tmp_path / 'f.csv')])


# --- the shortcut: the original intensities already written down ------------

def add_orig_column(path, original, blank_last=0):
    """Put the true original intensities into an `Iorig` column.

    `blank_last` rows are left empty, standing for lines whose paper the
    column does not cover; those must still go the long way round.
    """
    wb = openpyxl.load_workbook(path)
    ws = wb['Sheet1']
    ws.cell(1, ws.max_column + 1).value = 'Iorig'
    c = ws.max_column
    for k in range(len(original) - blank_last):
        ws.cell(k + 2, c).value = float(original[k])
    wb.save(path)


def test_a_written_down_original_is_used_as_it_stands(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    # a column that disagrees with what undoing the old correction would give
    add_orig_column(path, original * 2.0)
    app.main(['--lines', path, '--old', p_old, '--new', p_new,
              '--orig-col', 'Iorig', '--scale', str(SCALE),
              '--out-csv', str(tmp_path / 'g.csv')])
    lam = 1e8 / wn
    want = SCALE * original * 2.0 * np.exp(
        cal.apply_correction(cal.read_correction(p_new), lam))
    assert np.allclose(read_column(path), want, rtol=1e-9)


def test_rows_without_one_are_still_restored(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    add_orig_column(path, original, blank_last=5)
    app.main(['--lines', path, '--old', p_old, '--new', p_new,
              '--orig-col', 'Iorig', '--scale', str(SCALE),
              '--out-csv', str(tmp_path / 'h.csv')])
    lam = 1e8 / wn
    want = SCALE * original * np.exp(
        cal.apply_correction(cal.read_correction(p_new), lam))
    assert np.allclose(read_column(path), want, rtol=1e-9)


def test_auto_scale_is_refused_when_there_is_nothing_to_guess_from(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    add_orig_column(path, original)
    with pytest.raises(SystemExit):
        app.main(['--lines', path, '--old', p_old, '--new', p_new,
                  '--orig-col', 'Iorig',
                  '--out-csv', str(tmp_path / 'i.csv')])


def test_a_missing_orig_column_is_an_error(tmp_path):
    path, p_old, p_new, original, wn = make_case(tmp_path)
    with pytest.raises(SystemExit):
        app.main(['--lines', path, '--old', p_old, '--new', p_new,
                  '--orig-col', 'nosuch', '--scale', str(SCALE),
                  '--out-csv', str(tmp_path / 'j.csv')])
