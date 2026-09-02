"""Tests for `tools/attach_sugar_intensities.py`.

Run from the LineClass directory:  python -m pytest tests -q

The script copies the intensities printed in a source paper into a column of
the working line list, matching lines by their reconciled observed wavenumber
(1e8 divided by the vacuum wavelength in angstroms, in reciprocal centimetres).
Small workbooks are built here for it to work on.
"""
import os
import sys

import openpyxl
import pytest

TOOLS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     'tools')
sys.path.insert(0, TOOLS)

import attach_sugar_intensities as att      # noqa: E402


def make_case(tmp_path):
    """A source workbook and a line list built from part of it.

    The source has ten lines; the line list takes eight of them, one of them
    twice (an observed line with two candidate identifications), and adds one
    line whose wavenumber is in no source at all.  One of the eight is credited
    in the `ref` column to somebody else, because `ref` says whose
    classification the row carries, not who measured the line: it must be
    matched all the same.  The wavenumbers in the line list are the source ones
    perturbed by 1e-11 cm^-1, the size of a float round-trip.
    """
    src = str(tmp_path / 'source.xlsm')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Table 1'
    ws.append(['Wl', 'junk', 'Intensity', 'more', 'wn\nmean'])
    wns = [10000.0 + 137.0 * k for k in range(10)]
    ints = [1, 3, 20, 500, 9000, 2, 7, 40, 100, 6]
    for k in range(10):
        ws.append([1e8 / wns[k], 'x', ints[k], 'y', wns[k]])
    wb.save(src)

    lines = str(tmp_path / 'lines.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['own', 'Icor', 'ref', 'note'])
    taken = [0, 1, 2, 3, 3, 5, 6, 8, 9]       # index 3 appears twice
    for n, k in enumerate(taken):
        ref = 'Somebody else' if n == 6 else 'Paper'
        ws.append([wns[k] + 1e-11, 1.0, ref, 'x'])
    ws.append([99999.0, 1.0, 'Elsewhere', 'x'])
    wb.save(lines)
    return src, lines, [ints[k] for k in taken]


def run(lines, src, *extra):
    return att.main(['--lines', lines, '--sheet', 'Sheet1',
                     '--col-wn', 'own', '--col-ref', 'ref',
                     '--col-out', 'Iorig',
                     '--source', 'Paper=%s:Table 1:wn mean:Intensity' % src,
                     '--out-csv', lines + '.csv'] + list(extra))


def column(path, name):
    ws = openpyxl.load_workbook(path)['Sheet1']
    header = [c.value for c in ws[1]]
    j = header.index(name)
    return [r[j] for r in ws.iter_rows(min_row=2, values_only=True)]


def test_values_are_copied(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    got = column(lines, 'Iorig')
    assert got[:len(want)] == [float(v) for v in want]


def test_a_wavenumber_in_no_source_is_left_empty(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    assert column(lines, 'Iorig')[-1] is None


def test_the_ref_column_does_not_decide_the_match(tmp_path):
    """A row credited to another author is still matched on its wavenumber."""
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    assert column(lines, 'ref')[6] == 'Somebody else'
    assert column(lines, 'Iorig')[6] == float(want[6])


def test_one_source_line_may_serve_two_rows(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    got = column(lines, 'Iorig')
    assert got[3] == got[4] == 500.0


def test_nothing_else_changes(tmp_path):
    src, lines, want = make_case(tmp_path)
    before = column(lines, 'own'), column(lines, 'ref')
    run(lines, src)
    assert (column(lines, 'own'), column(lines, 'ref')) == before


def test_a_backup_keeps_the_workbook_as_it_was(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    bak = os.path.splitext(lines)[0] + '.bak_iorig.xlsx'
    assert 'Iorig' not in [c.value for c in openpyxl.load_workbook(bak)['Sheet1'][1]]


def test_dry_run_writes_nothing(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src, '--dry-run')
    assert 'Iorig' not in [c.value for c in openpyxl.load_workbook(lines)['Sheet1'][1]]


def test_a_second_run_updates_the_column_in_place(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    run(lines, src, '--no-backup')
    header = [c.value for c in openpyxl.load_workbook(lines)['Sheet1'][1]]
    assert header.count('Iorig') == 1
    assert column(lines, 'Iorig')[0] == float(want[0])


def test_a_value_already_there_survives_a_run_that_finds_no_source(tmp_path):
    """Rows matched by no source keep whatever the column already held."""
    src, lines, want = make_case(tmp_path)
    run(lines, src)
    run(lines, src, '--no-backup', '--tol', '1e-12')
    assert column(lines, 'Iorig')[0] == float(want[0])


def test_a_wavenumber_too_far_away_is_not_matched(tmp_path):
    src, lines, want = make_case(tmp_path)
    run(lines, src, '--tol', '1e-12')
    assert set(column(lines, 'Iorig')) == {None}


def test_two_sources_are_both_searched(tmp_path):
    """With the list split over two files, each row finds its own file."""
    src, lines, want = make_case(tmp_path)
    other = str(tmp_path / 'other.xlsm')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Table 1'
    ws.append(['Intensity', 'wn mean'])
    ws.append([777, 99999.0])
    wb.save(other)
    att.main(['--lines', lines, '--sheet', 'Sheet1', '--col-wn', 'own',
              '--col-ref', 'ref', '--col-out', 'Iorig',
              '--source', 'Paper=%s:Table 1:wn mean:Intensity' % src,
              '--source', 'Other=%s:Table 1:wn mean:Intensity' % other,
              '--out-csv', lines + '.csv'])
    got = column(lines, 'Iorig')
    assert got[0] == float(want[0]) and got[-1] == 777.0


def test_columns_may_be_named_by_number(tmp_path):
    src, lines, want = make_case(tmp_path)
    att.main(['--lines', lines, '--sheet', 'Sheet1', '--col-wn', 'own',
              '--col-ref', 'ref', '--col-out', 'Iorig',
              '--source', 'Paper=%s:Table 1:5:3' % src,
              '--out-csv', lines + '.csv'])
    assert column(lines, 'Iorig')[0] == 1.0


def test_a_missing_column_is_refused(tmp_path):
    src, lines, want = make_case(tmp_path)
    with pytest.raises(SystemExit):
        att.main(['--lines', lines, '--col-wn', 'nosuch',
                  '--source', 'Paper=%s:Table 1:5:3' % src])
