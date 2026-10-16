"""Tests for levels_pub.py: the article's level table filled from Set 3.

Run from the LineClass directory:  python -m pytest tests -q

The tests fix:

  * the join of a row to the calculated level list: by level_id first, then
    by parity, J and the nearest calculated energy, a near tie decided by
    the lid order of the neighboring rows;
  * the columns from LOPT, levels_set3.csv and the set's IDEN2 table, and
    the refusal of a levels_set3.csv that belongs to another fit or has no
    systematic parts;
  * the N/R/D notes and the Q added after a comma;
  * a row inserted for a level beyond the table, and a rerun that replaces
    the columns instead of adding more;
  * a column added to a workbook written before it existed (iden2_id after
    lid), with the workbook's own formulas following the cells they name.
"""
import math
import os
import sys
import types

import openpyxl
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import levels_pub as L                     # noqa: E402

ID = '059003.%06d'.__mod__


def calc_level(lid, parity, J, E, conf='4f3', term='4I'):
    return dict(lid=lid, parity=parity, J=J, E_calc=E, conf=conf, term=term)


CALC = [calc_level(1, 'odd', 4.5, 82.0),
        calc_level(2, 'odd', 5.5, 1473.2),
        calc_level(3, 'odd', 4.5, 9000.0),
        calc_level(4, 'odd', 4.5, 9001.0),
        calc_level(5, 'odd', 2.5, 150000.0, conf='4f5d2', term='2D'),
        calc_level(10, 'even', 1.5, 60000.0)]
LID_OF_ID = {ID(1): 1, ID(31): 2, ID(40): 3, ID(41): 4, ID(90): 5,
             ID(50): 10}


def row(n, parity, J, Ec, level_id=''):
    return dict(row=n, parity=parity, J=J, Ec=Ec, energy=None,
                level_id=level_id)


# ---------------------------------------------------------------------------
# The join
# ---------------------------------------------------------------------------
def test_a_row_with_a_level_id_is_joined_by_it_whatever_its_Ec():
    lids, problems = L.match_rows([row(2, 'o', 5.5, 1470.0, ID(31))],
                                  CALC, LID_OF_ID)
    assert lids == [2] and problems == []


def test_a_row_without_one_takes_the_nearest_level_of_its_parity_and_J():
    rows = [row(2, 'o', 4.5, 83.5), row(3, 'e', 1.5, 60002.0)]
    lids, problems = L.match_rows(rows, CALC, LID_OF_ID)
    assert lids == [1, 10] and problems == []


def test_a_near_tie_is_decided_by_the_order_of_the_rows():
    # 9000.5 is as near to lid 3 as to lid 4; the row below it holds lid 4
    rows = [row(2, 'o', 4.5, 82.0, ID(1)), row(3, 'o', 4.5, 9000.5),
            row(4, 'o', 4.5, 9001.0, ID(41))]
    lids, problems = L.match_rows(rows, CALC, LID_OF_ID)
    assert lids == [1, 3, 4] and problems == []


def test_a_wrong_parity_or_J_is_reported_not_joined():
    lids, problems = L.match_rows([row(2, 'e', 5.5, 1473.2, ID(31))],
                                  CALC, LID_OF_ID)
    assert lids == [None] and len(problems) == 1


# ---------------------------------------------------------------------------
# The values
# ---------------------------------------------------------------------------
def lopt(lid, E, n='12, 2L, 1D', D1='0.01', D2='0.02'):
    return dict(Designation=lid, Energy=E, D1=D1, D2stat=D2, N_lines=n)


def set3(lid, E, u_stat, u_cal, u_hfs, u_k, n='12', S='0.1'):
    tot = math.sqrt(u_stat ** 2 + u_cal ** 2 + u_hfs ** 2 + u_k ** 2)
    return {'level_id': lid, 'E_cg': E, 'u_stat': str(u_stat),
            'u_cal': str(u_cal), 'u_hfs': str(u_hfs), 'u_kappa': str(u_k),
            'D2tot*': '%.7f' % tot, 'N_lines': n, 'S': S}


def test_the_values_combine_the_parts_and_strip_the_L_and_D_counts():
    v = L.level_values([lopt(ID(31), '1398.35')],
                       [set3(ID(31), '1398.35', 0.02, 0.003, 0.004, 0.001)])
    E, u_stat, u_sys, u_tot, n, S = v[ID(31)]
    assert E == 1398.35 and n == 12
    assert u_sys == pytest.approx(math.sqrt(0.003 ** 2 + 0.004 ** 2 + 1e-6))
    assert u_tot == pytest.approx(math.hypot(0.02, u_sys))


def test_a_levels_set3_of_another_fit_is_refused():
    with pytest.raises(SystemExit):
        L.level_values([lopt(ID(31), '1398.40')],
                       [set3(ID(31), '1398.35', 0.02, 0.003, 0.004, 0.001)])


def test_a_levels_set3_without_systematic_parts_is_refused():
    s = set3(ID(31), '1398.35', 0.02, 0.003, 0.004, 0.001)
    s['u_cal'] = s['D2tot*'] = ''
    with pytest.raises(SystemExit):
        L.level_values([lopt(ID(31), '1398.35')], [s])


def test_the_notes_combine_their_letters():
    assert L.notes_of({'a', 'b'}, {'b'}, {'c'}) == {'a': 'N', 'b': 'NR',
                                                    'c': 'D'}


def test_Q_stands_alone_or_follows_a_comma():
    assert L.notes_of({'a', 'b'}, {'b'}, set(), {'b', 'q'}) == {
        'a': 'N', 'b': 'NR,Q', 'q': 'Q'}


def test_references_move_with_an_inserted_column_or_row():
    assert L.shift_refs('=U2-O2', col_at=20) == '=V2-O2'
    assert L.shift_refs('=SUM($U$2:U9)+T1', col_at=21) == '=SUM($V$2:V9)+T1'
    assert L.shift_refs('=U2-O5', row_at=5) == '=U2-O6'
    assert L.shift_refs('=MAX(U:U)', col_at=21) == '=MAX(V:V)'


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------
HEAD = ['parity', 'conf', 'term', 'energy', 'Ec', 'J', 'level_id']


def workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEAD)
    ws.append(['o', '4f3', '4I*', 0, 82.0, 4.5, ID(1)])
    ws.append(['o', '4f3', '4I*', 1398.3, 1473.1, 5.5, None])
    ws.append(['e', '4f2 5d', '2P', None, 60000.4, 1.5, None])
    wb.save(path)


def write_tsv(path, rows):
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        fh.write('\t'.join(rows[0]) + '\n')
        for r in rows:
            fh.write('\t'.join(str(r[k]) for k in rows[0]) + '\n')


def write_csv(path, rows):
    import csv
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


@pytest.fixture
def setup(tmp_path):
    pub = tmp_path / 'levels_pub.xlsx'
    workbook(pub)
    sd = tmp_path / 'final'
    sd.mkdir()
    write_tsv(sd / L.LOPT_LEVELS, [
        lopt(ID(1), '0.0000000', n='22, 3L', D1='0.0065', D2='0.0'),
        lopt(ID(31), '1398.3505782'), lopt(ID(50), '60010.1'),
        lopt(ID(90), '141300.5', n='4')])
    write_csv(sd / L.SET3_LEVELS, [
        set3(ID(1), '0.0000000', 0.0, 0.0, 0.016, 0.0, n='22', S='0.343'),
        set3(ID(31), '1398.3505782', 0.0077, 0.0027, 0.0155, 0.0004),
        set3(ID(50), '60010.1', 0.02, 0.01, 0.01, 0.0),
        set3(ID(90), '141300.5', 0.1, 0.05, 0.02, 0.01, n='4')])
    (sd / 'IDEN2').mkdir()
    (sd / L.ID_MAP).write_text(
        'level_id\tIDEN_id\n%s\t1253\n%s\t1252\n%s\t900\n%s\t658\n'
        % (ID(1), ID(31), ID(50), ID(90)), encoding='utf-8')
    quest = tmp_path / 'questionable_levels.txt'
    quest.write_text('level_id\tE\n%s\t141300.5\n%s\t60010.1\n'
                     % (ID(90), ID(50)), encoding='utf-8')
    new = tmp_path / 'new_levels.txt'
    new.write_text('level_id\tE\n%s\t141300.5\n' % ID(90), encoding='utf-8')
    rev = tmp_path / 'revised.csv'
    rev.write_text('level_id,E_input\n%s,1398.0\n' % ID(31), encoding='utf-8')
    disc = tmp_path / 'discarded.csv'
    disc.write_text('level_id\n', encoding='utf-8')
    cfg = types.SimpleNamespace(new_levels=str(new), level_overrides=str(rev),
                                discarded_levels=str(disc))
    cfg.questionable = str(quest)
    return pub, sd, cfg


def run(pub, sd, cfg, **kw):
    kw.setdefault('calc', CALC)
    kw.setdefault('lid_of_id', LID_OF_ID)
    return L.build(str(pub), str(sd), cfg, questionable_path=cfg.questionable,
                   **kw)


def read_back(path):
    ws = openpyxl.load_workbook(path).active
    head = [c.value for c in ws[1]]
    return head, [dict(zip(head, r)) for r in
                  ws.iter_rows(min_row=2, values_only=True)]


def test_the_table_is_completed_and_a_rerun_replaces_the_columns(setup):
    pub, sd, cfg = setup
    for _ in range(2):
        wb, ws, cells, report = run(pub, sd, cfg)
        L.write(wb, ws, cells, str(pub))
    head, rows = read_back(pub)
    assert head == HEAD + L.NEW_COLUMNS          # once, not twice
    by_id = {r['level_id']: r for r in rows}
    # filled by the join
    assert by_id[ID(31)]['lid'] == 2 and by_id[ID(31)]['Notes'] == 'R'
    assert by_id[ID(31)]['iden2_id'] == 1252
    assert by_id[ID(50)]['lid'] == 10 and by_id[ID(50)]['Notes'] == 'Q'
    assert by_id[ID(31)]['N_lines'] == 12
    assert by_id[ID(31)]['Eopt'] == pytest.approx(1398.3505782)
    assert by_id[ID(1)]['unc_stat'] == 0.0
    assert by_id[ID(1)]['unc_tot'] == pytest.approx(0.016)
    # the level above the table got a row, after the last odd one
    new = by_id[ID(90)]
    assert new['lid'] == 5 and new['Notes'] == 'N,Q' and new['Ec'] == 150000.0
    assert new['iden2_id'] == 658
    assert new['term'] == '2D*' and new['parity'] == 'o'
    assert [r['parity'] for r in rows] == ['o', 'o', 'o', 'e']
    assert len(rows) == 4                        # inserted once, not twice


def test_a_level_of_the_fit_with_no_possible_row_stops_the_run(setup):
    pub, sd, cfg = setup
    with pytest.raises(SystemExit):
        run(pub, sd, cfg, calc=CALC[:4] + CALC[5:],
            lid_of_id={k: v for k, v in LID_OF_ID.items() if v != 5})


def test_a_new_column_goes_after_lid_and_the_formulas_follow(setup):
    # the workbook as the seven-column version of the tool left it, with a
    # formula column of the user's own, dE_O-C = Eopt - Ec, before them
    pub, sd, cfg = setup
    old = [c for c in L.NEW_COLUMNS if c != 'iden2_id']
    wb = openpyxl.load_workbook(pub)
    ws = wb.active
    ws.insert_cols(6)
    ws.cell(1, 6, 'dE_O-C')
    for k, name in enumerate(old):
        ws.cell(1, len(HEAD) + 2 + k, name)
    eopt = openpyxl.utils.get_column_letter(len(HEAD) + 3)
    for r in (2, 3, 4):
        ws.cell(r, 6, '=%s%d-E%d' % (eopt, r, r))
    ws.auto_filter.ref = 'A1:%s4' % openpyxl.utils.get_column_letter(
        ws.max_column)
    wb.save(pub)
    wb, ws, cells, report = run(pub, sd, cfg)
    L.write(wb, ws, cells, str(pub))
    ws = openpyxl.load_workbook(pub).active
    head = [c.value for c in ws[1]]
    assert head[head.index('lid') + 1] == 'iden2_id'
    assert head[head.index('iden2_id') + 1] == 'Eopt'
    new_eopt = openpyxl.utils.get_column_letter(head.index('Eopt') + 1)
    assert ws.cell(2, 6).value == '=%s2-E2' % new_eopt
    assert ws.auto_filter.ref.endswith(
        openpyxl.utils.get_column_letter(len(head)) + '5')
    # the row inserted for lid 5 took the formula of the row above it
    rows = {ws.cell(r, head.index('level_id') + 1).value: r
            for r in range(2, ws.max_row + 1)}
    r = rows[ID(90)]
    assert ws.cell(r, 6).value == '=%s%d-E%d' % (new_eopt, r, r)
