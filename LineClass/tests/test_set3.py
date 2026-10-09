"""Tests for Set 3, the published set at the centers of gravity
(make_set3.py; hfs_correction.Model.cg_shift; config.derived_from).

Run from the LineClass directory:  python -m pytest tests -q

The tests fix:

  * the move of a line to its center of gravity: kappa * D per component,
    the flag and companion kappas of the head frame, the exception rows, a
    resolved level's S taken off whole, and the identity with the head
    frame's hfs_shift;
  * its uncertainty, and the part owed to levels without a determined A;
  * the rules of the derived set: merged companions and their light, the
    F = 2 partner of a resolved level's line, the characters;
  * the table: lines moved, fit uncertainties and dif_wn_O-C kept, energies
    at the centers of gravity;
  * that a derived set refuses the classification and the level tools.
"""
import csv
import math
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import config                              # noqa: E402
import hfs_correction as H                 # noqa: E402
import make_set3 as M3                     # noqa: E402

KAPPA = (('flag', (1.0, 0.0)), ('plain_1974', (0.944, 0.014)),
         ('plain_1969', (0.633, 0.035)), ('c', (0.20, 0.12)))

LOW, UPP, OTHER, RESOLVED, UNDET = ('059003.000100', '059003.000300',
                                    '059003.000301', '059003.000303',
                                    '059003.000304')
GROUND = M3.GROUND
EXC_HEADER = '\t'.join(H.EXCEPTION_COLUMNS) + '\n'
SAT_HEADER = '\t'.join(H.SATELLITE_COLUMNS) + '\n'


def a_table(tmp_path):
    path = tmp_path / 'A.csv'
    path.write_text(
        'level_id,cfg,J,A_cm-1,u_A,source,n_flagged\n'
        f'{GROUND},4f3,4.5,+0.0300,0.0020,flag interval,0\n'
        f'{LOW},f26s,3.5,-0.0400,0.0040,flag interval,0\n'
        f'{UPP},f26p,2.5,+0.1000,0.0100,flag interval,0\n'
        f'{OTHER},f5d2,1.5,+0.0200,0.0030,flag interval,0\n'
        f'{RESOLVED},f26s,0.5,+0.5000,0.0100,flag interval,0\n'
        f'{UNDET},f5d2,2.5,+0.0000,0.0500,not determined,0\n',
        encoding='utf-8', newline='\n')
    return str(path)


def tsv(tmp_path, name, header, rows):
    path = tmp_path / name
    path.write_text(header + ''.join('\t'.join(r) + '\n' for r in rows),
                    encoding='utf-8', newline='\n')
    return str(path)


def make_model(tmp_path, exceptions=(), satellites=()):
    return H.Model(config.HfsSettings(
        apply=True, A_levels=a_table(tmp_path), kappa=KAPPA,
        resolved_levels=(RESOLVED,),
        satellites=tsv(tmp_path, 'sat.txt', SAT_HEADER, satellites),
        kappa_exceptions=tsv(tmp_path, 'exc.txt', EXC_HEADER, exceptions)))


def S(lid):
    J, A = {GROUND: (4.5, 0.03), LOW: (3.5, -0.04), UPP: (2.5, 0.1),
            OTHER: (1.5, 0.02), RESOLVED: (0.5, 0.5), UNDET: (2.5, 0.0)}[lid]
    return 2.5 * A * J


# ---------------------------------------------------------------------------
# The move of a line
# ---------------------------------------------------------------------------
def test_a_plain_line_moves_by_kappa_D(tmp_path):
    m = make_model(tmp_path)
    D = S(UPP) - S(LOW)
    got = m.cg_shift('', 30000.0, [(LOW, UPP, 1.0)])
    assert got.shift == pytest.approx(-0.944 * D)
    assert got.D == pytest.approx(D) and got.R == 0.0
    assert got.kappa == pytest.approx(0.944)
    # the uncertainty: kappa * u_D and u_kappa * D in quadrature
    u_D = math.hypot(2.5 * 3.5 * 0.004, 2.5 * 2.5 * 0.01)
    assert got.u == pytest.approx(math.hypot(0.944 * u_D, 0.014 * D))
    assert got.u_undet == 0.0


def test_a_flagged_line_moves_to_head_minus_D(tmp_path):
    m = make_model(tmp_path)
    D = S(UPP) - S(LOW)
    got = m.cg_shift('*v', 30000.0, [(LOW, UPP, 1.0)])
    assert got.shift == pytest.approx(-D)
    assert got.kappa == 1.0


def test_a_flagged_blend_moves_its_components_by_their_kappas(tmp_path):
    m = make_model(tmp_path)
    d1, d2 = S(UPP) - S(LOW), S(UPP) - S(OTHER)
    got = m.cg_shift('*r', 30000.0, [(LOW, UPP, 0.7), (OTHER, UPP, 0.3)])
    # the strongest component takes the flag's kappa, the other the plain
    assert got.shift == pytest.approx(-(0.7 * 1.0 * d1 + 0.3 * 0.944 * d2))
    assert got.kappa == pytest.approx(0.7 + 0.3 * 0.944)


def test_the_main_line_of_companions_moves_to_head_minus_D(tmp_path):
    m = make_model(tmp_path, satellites=[
        ('29999.7', '30000', LOW, UPP, '1', '', '', '')])
    D = S(UPP) - S(LOW)
    assert m.cg_shift('', 30000.0, [(LOW, UPP, 1.0)]).shift == \
        pytest.approx(-D)


@pytest.mark.parametrize('cls, kappa', [('head', 1.0), ('cg', 0.0),
                                        ('c', 0.20)])
def test_an_exception_row_sets_the_kappa(tmp_path, cls, kappa):
    m = make_model(tmp_path, exceptions=[
        ('30000.0', LOW, UPP, cls, '', '', '')])
    D = S(UPP) - S(LOW)
    assert m.cg_shift('', 30000.0, [(LOW, UPP, 1.0)]).shift == \
        pytest.approx(-kappa * D)


def test_a_resolved_level_s_S_is_taken_off_whole(tmp_path):
    m = make_model(tmp_path)
    got = m.cg_shift('', 30000.0, [(RESOLVED, UPP, 1.0)])
    # the line ends on the F = I + J sublevel: S(resolved) leaves at kappa 1
    assert got.D == pytest.approx(S(UPP))
    assert got.R == pytest.approx(-S(RESOLVED))
    assert got.shift == pytest.approx(-(0.944 * S(UPP) - S(RESOLVED)))
    assert got.u == pytest.approx(math.hypot(
        0.944 * 2.5 * 2.5 * 0.01 + 2.5 * 0.5 * 0.01, 0.014 * S(UPP)))


def test_an_unresolved_row_of_a_resolved_level_is_not_moved(tmp_path):
    m = make_model(tmp_path, exceptions=[
        ('30000.0', RESOLVED, UPP, 'cg', RESOLVED, '', '')])
    got = m.cg_shift('', 30000.0, [(RESOLVED, UPP, 1.0)])
    assert got.shift == 0.0 and got.R == 0.0


@pytest.mark.parametrize('char, comps, exc, sat', [
    ('', [(LOW, UPP, 1.0)], (), ()),
    ('*r', [(LOW, UPP, 0.6), (OTHER, UPP, 0.4)], (), ()),
    ('c', [(RESOLVED, UPP, 1.0)], (), ()),
    ('', [(RESOLVED, UPP, 1.0)], [('30000', RESOLVED, UPP, 'cg', RESOLVED,
                                   '', '')], ()),
    ('w', [(LOW, UPP, 0.5), (OTHER, UPP, 0.5)],
     [('30000', OTHER, UPP, 'partial:0-1', '', '', '')], ()),
    ('', [(LOW, UPP, 1.0)], (), [('29999.7', '30000', LOW, UPP, '1', '',
                                  '', '')]),
])
def test_the_identity_with_the_head_frame(tmp_path, char, comps, exc, sat):
    m = make_model(tmp_path, exceptions=exc, satellites=sat)
    head = m.line_shift(char, 30000.0, comps)
    cg = m.cg_shift(char, 30000.0, comps)
    assert head[0] - cg.shift == pytest.approx(cg.D + cg.R, abs=1e-12)
    assert cg.kappa == pytest.approx(head[2])


def test_the_undetermined_part_of_the_uncertainty(tmp_path):
    m = make_model(tmp_path)
    got = m.cg_shift('', 30000.0, [(LOW, UNDET, 1.0)])
    assert got.shift == pytest.approx(-0.944 * -S(LOW))   # S(UNDET) = 0
    u_low, u_und = 2.5 * 3.5 * 0.004, 2.5 * 2.5 * 0.05
    full = math.hypot(0.944 * math.hypot(u_low, u_und), 0.014 * -S(LOW))
    known = math.hypot(0.944 * u_low, 0.014 * -S(LOW))
    assert got.u == pytest.approx(full)
    assert got.u_undet == pytest.approx(math.sqrt(full ** 2 - known ** 2))


def test_a_line_without_an_accepted_transition_stays(tmp_path):
    got = make_model(tmp_path).cg_shift('*r', 30000.0, [])
    assert (got.shift, got.u, got.D, got.R) == (0.0, 0.0, 0.0, 0.0)


# ---------------------------------------------------------------------------
# The rules of the derived set
# ---------------------------------------------------------------------------
def row(wn_key, low='', upp='', accepted='', bf='', char='', intens='100',
        wn=None, hfs_shift='0.0', from_comp='0.0', unc='0.02', grade='2A'):
    return {'wn_obs': repr(wn if wn is not None else wn_key),
            'wn_key': repr(wn_key), 'unc_wn_obs': unc,
            'obs_intens': intens, 'char': char, 'low_id': low,
            'upp_id': upp, 'grade': grade, 'accepted': accepted,
            'BF': bf, 'hfs_shift': hfs_shift, 'u_hfs_shift': '0.0123',
            'u_hfs_undet': '0.0', 'obs_intens_hfs': from_comp,
            'dif_wn_O-C': '', 'low_E': '', 'upp_E': '', 'rwn': ''}


def lines_of(*rows, model=None):
    """The lines of `rows`; with `model`, each given the head-frame shift
    that model writes into a table, as classify_lines.py would."""
    out = {}
    for r in rows:
        out.setdefault(M3.key4(r['wn_key']), []).append(r)
    lines = {k: M3.Line(k, v) for k, v in out.items()}
    if model is not None:
        for ln in lines.values():
            ln.hfs_shift = model.line_shift(ln.char, ln.wn_key,
                                            ln.components)[0]
    return lines


def test_companions_are_merged_and_a_blended_one_keeps_its_rest(tmp_path):
    m = make_model(tmp_path, satellites=[
        ('29999.7', '30000', LOW, UPP, '1', '', '', ''),
        ('29999.4', '30000', LOW, UPP, '2', '1', '', ''),
        ('25000.1', '', OTHER, UPP, '1', '', '', '')])
    lines = lines_of(
        row(30000.0, LOW, UPP, '1', '1.0', intens='170', from_comp='70'),
        row(29999.7, LOW, UPP, '', '', intens='40', grade='hfs'),
        row(29999.4, OTHER, UPP, '1', '1.0', intens='50'),
        row(25000.1, OTHER, UPP, '', '', intens='5', grade='hfs'), model=m)
    report = {}
    M3.move_lines(lines, m, report)
    M3.merge_companions(lines, m, report)
    main, whole, part, headless = (lines[M3.key4(k)] for k in
                                   (30000.0, 29999.7, 29999.4, 25000.1))
    assert whole.merged_into == 30000.0 and whole.intens3 == 0.0
    assert part.merged_into is None
    assert part.intens3 == pytest.approx(50 - (70 - 40))
    assert main.intens3 == pytest.approx(170)          # already summed
    assert headless.char3 == 'hfs' and headless.merged_into is None


def test_the_partner_of_a_resolved_level_s_line_is_merged(tmp_path):
    m = make_model(tmp_path)
    J, A, u_A, _ = m.levels[RESOLVED]
    delta = A * (2.5 + J)       # the lower level: the partner lies above
    lines = lines_of(row(30000.0, RESOLVED, UPP, '1', '1.0', intens='70'),
                     row(30000.0 + delta + 0.01, intens='50'),
                     row(30000.0 + delta + 5.0, intens='9'), model=m)
    listed = [(l.wn, l.unc, l.key) for l in lines.values()]
    report = {}
    M3.move_lines(lines, m, report)
    M3.pair_resolved(lines, m, listed, report)
    main = lines[M3.key4(30000.0)]
    partner = lines[M3.key4(30000.0 + delta + 0.01)]
    assert partner.merged_into == 30000.0 and partner.intens3 == 0.0
    assert main.intens3 == pytest.approx(120) and main.has_partner
    assert len(report['pairs']) == 1
    assert lines[M3.key4(30000.0 + delta + 5.0)].merged_into is None


def test_an_accepted_line_is_never_taken_as_a_partner(tmp_path):
    m = make_model(tmp_path)
    J, A, _, _ = m.levels[RESOLVED]
    delta = A * (2.5 + J)
    lines = lines_of(row(30000.0, RESOLVED, UPP, '1', '1.0'),
                     row(30000.0 + delta, LOW, OTHER, '1', '1.0'), model=m)
    report = {}
    M3.move_lines(lines, m, report)
    M3.pair_resolved(lines, m, [(l.wn, l.unc, l.key)
                                for l in lines.values()], report)
    assert 'pairs' not in report and len(report['unpaired']) == 1


def test_characters(tmp_path, monkeypatch):
    monkeypatch.setattr(M3, 'fwhm_cm', lambda wn: 0.7)
    m = make_model(tmp_path)
    D = S(UPP) - S(LOW)                  # +0.975: shaded to longer waves
    lines = lines_of(row(30000.0, LOW, UPP, '1', '1.0', char='*r'),
                     row(30001.0, LOW, UPP, '1', '1.0', char='*v'),
                     row(30002.0, LOW, UPP, '1', '1.0', char='c'),
                     row(30003.0, OTHER, UPP, '1', '1.0', char='w'),
                     row(30004.0, LOW, UPP, '1', '1.0', char=''),
                     row(30005.0, char='c'), model=m)
    report = {}
    M3.move_lines(lines, m, report)
    M3.characters(lines, m, report)
    got = {k: lines[M3.key4(k)].char3 for k in
           (30000.0, 30001.0, 30002.0, 30003.0, 30004.0, 30005.0)}
    assert D > 0.7 > S(UPP) - S(OTHER)
    assert got == {30000.0: 'hfs,l', 30001.0: 'hfs,s', 30002.0: 'hfs,l',
                   30003.0: 'w', 30004.0: '', 30005.0: 'c'}
    # the *v line's D says the pattern is shaded the other way
    assert [ln.wn_key for ln, _ in report['flag_sign']] == [30001.0]


def test_the_table_is_moved_and_keeps_the_fit_s_uncertainties(tmp_path):
    m = make_model(tmp_path)
    D = S(UPP) - S(LOW)
    r1 = row(30000.0, LOW, UPP, '1', '1.0', hfs_shift=repr(0.056 * D))
    r1.update({'low_E': '1000.0', 'upp_E': '31000.1', 'rwn': '30000.1',
               'dif_wn_O-C': repr(30000.0 - (30000.1 - 0.056 * D))})
    r2 = row(20000.0, intens='3')
    fields = list(r1)
    lines = lines_of(r1, r2)
    report = {}
    M3.move_lines(lines, m, report)
    out_csv, out_xlsx = str(tmp_path / 't.csv'), str(tmp_path / 't.xlsx')
    assert M3.write_table(fields, [r1, r2], lines, m, out_csv, out_xlsx,
                          report) == 2
    with open(out_csv, encoding='utf-8', newline='') as fh:
        got = list(csv.DictReader(fh))
    s_g = S(GROUND)
    assert float(got[0]['wn_obs']) == pytest.approx(30000.0 - 0.944 * D,
                                                    abs=5e-5)
    assert got[0]['wn_measured'] == repr(30000.0)
    assert float(got[0]['hfs_shift']) == 0.0
    assert got[0]['u_hfs_shift'] == '0.0123'          # the fit's, kept
    assert float(got[0]['low_E']) == pytest.approx(1000.0 - S(LOW) + s_g)
    assert float(got[0]['upp_E']) == pytest.approx(31000.1 - S(UPP) + s_g)
    # a single line: the classification's dif is wn_obs - rwn exactly
    assert report['dif_check'] < 1e-4
    assert got[1]['wn_obs'] == repr(20000.0)          # not moved
    assert os.path.exists(out_xlsx)


def test_two_lines_moved_onto_one_wavenumber_stop_the_run(tmp_path):
    m = make_model(tmp_path)
    D = S(UPP) - S(LOW)
    r1 = row(30000.0, LOW, UPP, '1', '1.0', char='*r')
    r2 = row(round(30000.0 - D, 4))
    lines = lines_of(r1, r2)
    M3.move_lines(lines, m, {})
    with pytest.raises(SystemExit, match='both come to'):
        M3.write_table(list(r1), [r1, r2], lines, m, str(tmp_path / 't.csv'),
                       str(tmp_path / 't.xlsx'), {})


def test_a_table_the_model_did_not_write_stops_the_run(tmp_path):
    m = make_model(tmp_path)
    lines = lines_of(row(30000.0, LOW, UPP, '1', '1.0', hfs_shift='0.5'))
    with pytest.raises(SystemExit, match='disagrees with the model'):
        M3.move_lines(lines, m, {})


# ---------------------------------------------------------------------------
# A derived set refuses the tools that write assignments
# ---------------------------------------------------------------------------
def derived_set(tmp_path):
    d = tmp_path / 'final'
    d.mkdir()
    (d / config.CONFIG_NAME).write_text('derived_from = "make_set3.py"\n',
                                        encoding='utf-8')
    return d


def test_a_derived_set_is_named(tmp_path):
    d = derived_set(tmp_path)
    assert config.derived_from(str(d)) == 'make_set3.py'
    other = tmp_path / 'iter'
    other.mkdir()
    (other / config.CONFIG_NAME).write_text('', encoding='utf-8')
    assert config.derived_from(str(other)) == ''


def test_the_classification_refuses_a_derived_set(tmp_path):
    d = derived_set(tmp_path)
    with pytest.raises(SystemExit, match='derived set'):
        config.require_own_output(str(d / config.CONFIG_NAME),
                                  [str(d / 'line_classifications.csv')],
                                  'classify_lines.py')


def test_the_level_tools_refuse_a_derived_set(tmp_path):
    import insert_new_level as INL
    d = derived_set(tmp_path)
    with pytest.raises(SystemExit, match='derived set'):
        INL.require_unlocked([str(d / 'LOPT_input_lines.txt')],
                             'insert_new_level.py', unlock=True)


def test_derived_from_is_not_inherited(tmp_path):
    d = derived_set(tmp_path)
    child = tmp_path / 'child.toml'
    child.write_text('inherit = "final/%s"\n' % config.CONFIG_NAME,
                     encoding='utf-8')
    raw = config._read(str(child))
    assert 'derived_from' not in raw


# ---------------------------------------------------------------------------
# Merged lines are no lines of Set 3
# ---------------------------------------------------------------------------
def merged_lines(tmp_path):
    m = make_model(tmp_path, satellites=[
        ('29999.7', '30000', LOW, UPP, '1', '', '', '')])
    lines = lines_of(
        row(30000.0, LOW, UPP, '1', '1.0', intens='140', from_comp='40'),
        row(29999.7, LOW, UPP, '', '', intens='40', grade='hfs'),
        row(25000.0, OTHER, UPP, '1', '1.0', intens='7'), model=m)
    report = {}
    M3.move_lines(lines, m, report)
    M3.merge_companions(lines, m, report)
    return m, lines, report


def test_the_table_leaves_out_a_merged_line(tmp_path):
    m, lines, report = merged_lines(tmp_path)
    rows = [r for ln in lines.values() for r in ln.rows]
    out_csv = str(tmp_path / 't.csv')
    assert M3.write_table(list(rows[0]), rows, lines, m, out_csv,
                          str(tmp_path / 't.xlsx'), report) == 2
    with open(out_csv, encoding='utf-8', newline='') as fh:
        keys = [float(r['wn_key']) for r in csv.DictReader(fh)]
    assert keys == [30000.0, 25000.0] and report['rows_dropped'] == 1


def test_the_list_puts_a_merged_line_on_its_own_sheet(tmp_path):
    import openpyxl
    m, lines, report = merged_lines(tmp_path)
    header = ['own', 'unc_own', 'Icor', 'Ch.', 'own_corr', 'unc_own_corr']
    rows = [[30000.0, 0.02, 100.0, '*r', 30000.01, 0.02],
            [29999.7, 0.02, 40.0, None, 29999.71, 0.02],
            [25000.0, 0.02, 7.0, None, 25000.01, 0.02]]
    cols = {'wn_key': 'own', 'wn': 'own_corr', 'u_wn': 'unc_own_corr',
            'character': 'Ch.', 'intensity': 'Icor'}
    path = str(tmp_path / 'list.xlsx')
    M3.write_list(header, rows, lines, cols, path)
    wb = openpyxl.load_workbook(path)
    main = [r[0] for r in wb['Sheet1'].iter_rows(min_row=2, values_only=True)]
    gone = list(wb[M3.MERGED_SHEET].iter_rows(min_row=2, values_only=True))
    assert main == [30000.0, 25000.0]
    assert [g[0] for g in gone] == [29999.7]
    assert gone[0][-1] == 30000.0                   # merged_into


def iden2_dir(tmp_path, identified=()):
    d = tmp_path / 'IDEN2'
    d.mkdir()
    (d / 'dlv.dat').write_bytes(
        b'   80     30000.010      3333.3322  /        *r/       0.0020    11\r\n'
        b'   70     29999.710      3333.3655  /          /       0.0020    12\r\n'
        b'   50     25000.010      3999.9984  /          /       0.0030    13\r\n')
    (d / 'dlv_keys.txt').write_text(
        'line\twn_key\twn_written\n11\t30000.0\t30000.010\n'
        '12\t29999.7\t29999.710\n13\t25000.0\t25000.010\n',
        encoding='utf-8', newline='\n')
    (d / 'dlv_shown.txt').write_text(
        'line\twn_key\twn_obs\tu_obs\twn_shown\tu_shown\n'
        '12\t29999.7\t29999.71\t0.02\t29999.8\t0.03\n',
        encoding='utf-8', newline='\n')
    trans = ''.join(' ' * 60 + '%6d\n' % n for n in identified)
    (d / 'trans.dat').write_text(trans or 'no identifications\n',
                                 encoding='utf-8', newline='\n')
    return d


def test_iden2_loses_the_merged_line_s_row_and_keeps_the_numbers(tmp_path):
    m, lines, report = merged_lines(tmp_path)
    d = iden2_dir(tmp_path, identified=(11, 13))
    M3.remove_merged_from_iden2(str(d), lines, report)
    assert report['iden2_removed'] == [(12, 29999.7)]
    dlv = (d / 'dlv.dat').read_bytes().splitlines(keepends=True)
    assert [r.split()[-1] for r in dlv] == [b'11', b'13']
    assert all(r.endswith(b'\r\n') for r in dlv)    # IDEN2's own endings
    assert '\n12\t' not in (d / 'dlv_keys.txt').read_text(encoding='utf-8')
    assert (d / 'dlv_shown.txt').read_text(encoding='utf-8').count('\n') == 1
    # a second run finds nothing left to do
    M3.remove_merged_from_iden2(str(d), lines, report)
    assert report['iden2_removed'] == []


def test_an_identified_merged_line_stops_the_run(tmp_path):
    m, lines, report = merged_lines(tmp_path)
    d = iden2_dir(tmp_path, identified=(11, 12))
    before = (d / 'dlv.dat').read_bytes()
    with pytest.raises(SystemExit, match='identifies the merged'):
        M3.remove_merged_from_iden2(str(d), lines, report)
    assert (d / 'dlv.dat').read_bytes() == before
