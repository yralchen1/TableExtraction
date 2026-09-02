#!/usr/bin/env python
"""Put a new intensity-scale correction into the observed line list.

What this does, and why it is not just a multiplication
-------------------------------------------------------
The intensity column of the line list (`Icor` of `Pr3_lines.xlsx`) does not
hold the numbers the original author reported.  It holds them after an
intensity-scale correction has been applied to them:

    Icor(lambda) = K * I_original * exp( P_old(lambda) )

where `lambda` is the vacuum wavelength of the line in angstroms
(`lambda = 1e8/wn`, `wn` = wavenumber in cm^-1), `P_old` is the piecewise
polynomial of the correction file that was in force, and `K` is an overall
scale factor the earlier work carried along (in this project, 1000).  The
correction exists because the plate, the grating and the optics all respond
differently at different wavelengths, so two equally bright lines recorded in
different parts of the spectrum were written down as different numbers.

To adopt a NEW correction one therefore has to undo the old one first,
recovering the original intensities, and only then apply the new one:

    I_original = Icor / ( K * exp(P_old(lambda)) )
    Icor_new   = K * I_original * exp( P_new(lambda) )

The original intensities are whole numbers (they were read off the plate as
such), so the restored values are snapped to whole numbers by default; how far
they had to move is reported, and is the check that the restoration is right.

When the original intensities are known exactly
-----------------------------------------------
Restoring them is only necessary while they are not written down anywhere.  If
the line list carries a column holding the numbers the author actually printed
- `Iorig`, put there by `tools/attach_sugar_intensities.py` - then give
`--orig-col Iorig` and that column is used directly, with no undoing and no
snapping, so the new intensity column is exactly

    Icor_new = K * Iorig * exp( P_new(lambda) )

Rows whose `--orig-col` cell is empty (lines the column does not cover) still
go the long way round, through `--old`.  Because the scale factor K can then no
longer be guessed from the restored values - and cannot be guessed at all when
every line has its original - give it explicitly: `--scale 1000` keeps the
scale this project has always used.

Usage (from the LineClass directory):

    python tools/apply_calibration.py --old intensity_correction_functions.txt \
                                      --new intensity_correction_auto.txt
    python tools/apply_calibration.py --old ... --new ... --dry-run
    python tools/apply_calibration.py --orig-col Iorig --scale 1000 \
                                      --old ... --new ...

The workbook is copied to `<name>.bak.xlsx` before it is changed, unless
--no-backup is given, and the restored original intensities are written out as
a csv so that the whole operation can be checked line by line.
"""
import argparse
import math
import os
import shutil
import sys

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from calibrate_intensities import (read_correction,        # noqa: E402
                                   apply_correction)


def parse_args(argv=None):
    root = os.path.dirname(HERE)
    p = argparse.ArgumentParser(
        description='Undo one intensity-scale correction and apply another.')
    p.add_argument('--lines', default=os.path.join(root, 'Pr3_lines.xlsx'))
    p.add_argument('--sheet', default='Sheet1')
    p.add_argument('--col-wn', default='own', help='observed wavenumber, cm^-1')
    p.add_argument('--col-intensity', default='Icor')
    p.add_argument('--orig-col', default=None, metavar='NAME',
                   help='column holding the original, uncorrected intensities; '
                        'where it has a value the old correction is not undone '
                        'but this number is used as it stands')
    p.add_argument('--old', required=True, metavar='FILE',
                   help='the correction currently applied to the column')
    p.add_argument('--new', required=True, metavar='FILE',
                   help='the correction to apply instead')
    p.add_argument('--scale', default='auto',
                   help='the overall factor K; "auto" (default) takes the '
                        'power of ten nearest the weakest restored line')
    p.add_argument('--no-round', dest='round', action='store_false',
                   help='do not snap the restored intensities to whole numbers')
    p.add_argument('--no-backup', dest='backup', action='store_false')
    p.add_argument('--out-csv', default=os.path.join(root,
                                                     'intensity_rescale.csv'),
                   help='per-line record of the change')
    p.add_argument('--dry-run', action='store_true',
                   help='report what would change, write nothing')
    return p.parse_args(argv)


def main(argv=None):
    a = parse_args(argv)
    wb = openpyxl.load_workbook(a.lines)
    ws = wb[a.sheet]
    header = [c.value for c in ws[1]]
    try:
        cw = header.index(a.col_wn) + 1
        ci = header.index(a.col_intensity) + 1
    except ValueError:
        raise SystemExit('%s: no column named %r or %r in row 1'
                         % (a.lines, a.col_wn, a.col_intensity))
    co = None
    if a.orig_col:
        if a.orig_col not in header:
            raise SystemExit('%s: no column named %r in row 1'
                             % (a.lines, a.orig_col))
        co = header.index(a.orig_col) + 1

    rows, wn, iold, given = [], [], [], []
    for r in range(2, ws.max_row + 1):
        w, v = ws.cell(r, cw).value, ws.cell(r, ci).value
        if w is None or v is None:
            continue
        w, v = float(w), float(v)
        if w <= 0 or v <= 0:
            continue
        g = float('nan')
        if co is not None:
            gv = ws.cell(r, co).value
            if isinstance(gv, (int, float)) and gv > 0:
                g = float(gv)
        rows.append(r)
        wn.append(w)
        iold.append(v)
        given.append(g)
    wn = np.asarray(wn, float)
    iold = np.asarray(iold, float)
    given = np.asarray(given, float)
    lam = 1e8 / wn
    print('%d lines with a wavenumber and a positive intensity, '
          '%.2f - %.2f A' % (len(rows), lam.min(), lam.max()))

    p_old = apply_correction(read_correction(a.old), lam)
    p_new = apply_correction(read_correction(a.new), lam)

    known = np.isfinite(given)
    if co is not None:
        print('%d of them carry their original intensity in %r; the other %d '
              'have it restored from %s'
              % (int(known.sum()), a.orig_col, int((~known).sum()),
                 os.path.basename(a.old)))

    restored = iold / np.exp(p_old)
    if a.scale == 'auto':
        if known.all():
            raise SystemExit('--scale auto has nothing to guess from when every '
                             'line has its original intensity; give --scale')
        K = 10.0 ** round(math.log10(float(restored[~known].min())))
    else:
        K = float(a.scale)
    restored = restored / K
    if not known.all():
        print('scale factor K = %g   ->  restored original intensities run '
              'from %.4g to %.4g' % (K, restored[~known].min(),
                                     restored[~known].max()))
        if a.round:
            snapped = np.round(restored)
            snapped[snapped < 1] = 1.0
            off = np.abs(restored[~known] / snapped[~known] - 1.0)
            print('snapped to whole numbers: largest departure %.1f%%, '
                  '99th percentile %.1f%%' % (100 * off.max(),
                                              100 * np.percentile(off, 99)))
            if off.max() > 0.35:
                print('WARNING: some restored values are far from a whole '
                      'number; --no-round may be safer here')
            restored = snapped
    else:
        print('scale factor K = %g' % K)
    orig = np.where(known, given, restored)
    print('original intensities in force run from %.4g to %.4g'
          % (orig.min(), orig.max()))

    inew = K * orig * np.exp(p_new)
    ratio = inew / iold
    print('new intensities run from %.4g to %.4g' % (inew.min(), inew.max()))
    print('change per line: median factor %.4g, 5th %.4g, 95th %.4g, '
          'largest %.4g' % (np.median(ratio), np.percentile(ratio, 5),
                            np.percentile(ratio, 95), ratio.max()))

    if not a.dry_run:
        with open(a.out_csv, 'w', encoding='utf-8') as fh:
            fh.write('row,wn,lambda,I_old,P_old,I_original,P_new,I_new\n')
            for k in range(len(rows)):
                fh.write('%d,%.4f,%.4f,%.6g,%.6f,%.6g,%.6f,%.6g\n'
                         % (rows[k], wn[k], lam[k], iold[k], p_old[k],
                            orig[k], p_new[k], inew[k]))
        print('per-line record written to %s' % a.out_csv)

        if a.backup:
            bak = os.path.splitext(a.lines)[0] + '.bak.xlsx'
            shutil.copyfile(a.lines, bak)
            print('the workbook as it was is kept in %s' % bak)
        for k, r in enumerate(rows):
            ws.cell(r, ci).value = float(inew[k])
        wb.save(a.lines)
        print('column %r of %s rewritten (%d values)'
              % (a.col_intensity, a.lines, len(rows)))
    else:
        print('dry run: nothing written')
    return 0


if __name__ == '__main__':
    sys.exit(main())
