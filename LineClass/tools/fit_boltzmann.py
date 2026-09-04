#!/usr/bin/env python
"""Refit the intensity model on a Boltzmann plot of the accepted lines.

What a Boltzmann plot is
------------------------
In a plasma in local thermodynamic equilibrium the number of atoms sitting in
an upper energy level falls off exponentially with the energy of that level,
as exp(-Eup/kT); Eup is the energy of the upper level of the transition in
cm^-1 and kT is an effective excitation temperature written as an energy in
the same unit.  The energy radiated in a line is that population multiplied by
the strength of the transition, so an observed intensity Iobs satisfies

    Iobs = C * gA * (rwn/1e8) * exp(-Eup/kT) ,                            (*)

where gA is the statistical weight of the upper level times the transition
probability per second (from the calculated-transition file), rwn is the Ritz
wavenumber of the line in cm^-1, and 1e8/rwn is its vacuum wavelength in
angstroms.  C absorbs everything constant across the spectrum (the number of
emitters, the solid angle, the plate response already removed by the intensity
calibration).

Taking logarithms turns (*) into a straight line:

    y = ln( Iobs * (1e8/rwn) / gA )  =  ln C  -  Eup/kT .

Plotting y against Eup - the Boltzmann plot - and fitting a straight line
therefore gives kT from the slope (kT = -1/slope) and C from the intercept
(C = exp(intercept)).  Every accepted identification contributes one point.

NOTE - this is a different model from the one the previous Icalc column
obeyed, which was Icalc = C*gA*exp(-Eup/kT)/rwn, i.e. proportional to the
wavelength rather than to the wavenumber.  The two differ by a factor rwn^2
across the spectrum.  Adopting (*) is the point of this script.

What the script does
--------------------
1. Reads the accepted rows (accepted = 1) of line_classifications.csv and
   takes obs_intens and rwn from them; takes gA and Eup for the same level
   pair from the calculated-transition file.  A pair absent from that file has
   no gA (its transition was too weak to be printed) and cannot enter a
   Boltzmann plot, so it is skipped.
2. Fits y against Eup by unweighted least squares -> first C, kT.
3. Computes the predicted intensity of every fitted point from (*) and drops
   the points with |ln(Icalc/Iobs)| > clip (default 2, i.e. disagreeing by
   more than a factor e^2 = 7.4 either way).
4. Refits on what survives -> the final C and kT.
5. With --write, recomputes the Icalc column of every row of the
   calculated-transition file from (*) with the final C and kT, using that
   file's own gA, Eup and rwn, and saves the workbook in place (all other
   columns, sheets and formatting untouched).  A copy of the workbook is put
   beside it first, under the name printed by the script.

Usage:
    python tools/fit_boltzmann.py                 # fit and report only
    python tools/fit_boltzmann.py --write         # ... and update Icalc.xlsx
    python tools/fit_boltzmann.py --plot fit.png  # ... and draw the plot
    python tools/fit_boltzmann.py --clip 2.5 --points points.csv
"""
import argparse
import os
import shutil
import sys

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config  # noqa: E402
import gA_imputation  # noqa: E402

# 1e8/rwn[cm^-1] is the vacuum wavelength of the transition in angstroms.
ANGSTROM_PER_CM = gA_imputation.ANGSTROM_PER_CM


def straight_line(x, y):
    """Least-squares fit y = a + b*x; returns (C, kT, a, b, rms residual).

    C = exp(a) and kT = -1/b are the two constants of the intensity model.
    """
    b, a = np.polyfit(x, y, 1)
    if b >= 0:
        raise SystemExit("the Boltzmann plot rises with Eup (slope >= 0): no "
                         "positive temperature describes these intensities")
    resid = y - (a + b * x)
    return float(np.exp(a)), float(-1.0 / b), float(a), float(b), \
        float(np.sqrt(np.mean(resid ** 2)))


def model_intensity(C, kT, gA, Eup, rwn):
    """Icalc = C * gA * (rwn/1e8) * exp(-Eup/kT), the model fitted here.

    Delegates to gA_imputation so that the fit and the rest of the pipeline
    cannot drift apart: there is one implementation of the model.
    """
    return gA_imputation.impute_intensity(gA, Eup, rwn, C, kT)


def read_icalc(cfg) -> pd.DataFrame:
    """The calculated-transition file, with the level ids kept as written.

    Read cell by cell with openpyxl rather than with pandas: the ids are text
    in the workbook ('059003.000042'), and pandas' type inference turns such a
    column into floats, which loses the leading zero and so matches nothing.
    """
    from openpyxl import load_workbook
    wb = load_workbook(cfg.icalc_file, read_only=True, data_only=True)
    ws = wb[cfg.icalc.sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else '' for h in next(rows)]
    col = cfg.icalc.columns
    want = {'low_id': col['id1'], 'upp_id': col['id2'], 'gA': col['gA'],
            'Eup': col['Eup'], 'rwn': col['rwn'], 'Icalc': col['Icalc']}
    missing = [v for v in want.values() if v not in header]
    if missing:
        raise SystemExit(f"{cfg.icalc_file}[{cfg.icalc.sheet}]: no column(s) "
                         f"named {missing}")
    pos = {k: header.index(v) for k, v in want.items()}
    data = {k: [] for k in want}
    for r in rows:
        if r[pos['low_id']] is None or r[pos['upp_id']] is None:
            continue
        for k, i in pos.items():
            data[k].append(r[i])
    df = pd.DataFrame(data)
    for k in ('low_id', 'upp_id'):
        df[k] = df[k].astype(str).str.strip()
    for k in ('gA', 'Eup', 'rwn', 'Icalc'):
        df[k] = pd.to_numeric(df[k], errors='coerce')
    return df


def load_points(cfg, lines_csv):
    """One row per accepted identification whose transition has a known gA."""
    lines = pd.read_csv(lines_csv, dtype={'low_id': str, 'upp_id': str})
    acc = lines[(lines['accepted'].fillna(0) == 1)
                & lines['low_id'].notna() & lines['upp_id'].notna()].copy()

    ic = read_icalc(cfg)[['low_id', 'upp_id', 'gA', 'Eup']]

    df = acc.merge(ic, on=['low_id', 'upp_id'], how='left')
    n_all = len(df)
    df = df[(df['gA'] > 0) & (df['obs_intens'] > 0) & (df['rwn'] > 0)
            & df['Eup'].notna()].copy()
    df['y'] = np.log(df['obs_intens'] * ANGSTROM_PER_CM / df['rwn'] / df['gA'])
    return df, n_all


def rewrite_icalc(cfg, C, kT):
    """Recompute the Icalc column of the workbook in place; return a summary.

    Formula cells are replaced by their value.  openpyxl writes no cached
    results for a formula, so a formula left in place would read back as an
    empty cell for every consumer that asks for values (the whole pipeline
    does) until Excel had opened and recalculated the file.  The one formula
    column of this workbook is u_ln = LN(1+u%gA/100), which is recomputed here
    from the same definition; any other formula found is replaced by the
    result Excel last cached for it, and named in the returned list.
    """
    from openpyxl import load_workbook
    backup = os.path.join(os.path.dirname(cfg.icalc_file),
                          os.path.splitext(os.path.basename(cfg.icalc_file))[0]
                          + '_before_boltzmann.xlsx')
    if not os.path.exists(backup):
        shutil.copy2(cfg.icalc_file, backup)
    wb = load_workbook(cfg.icalc_file)
    ws = wb[cfg.icalc.sheet]
    cached = load_workbook(cfg.icalc_file, read_only=True,
                           data_only=True)[cfg.icalc.sheet]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    col = cfg.icalc.columns
    need = [col['Icalc'], col['gA'], col['Eup'], col['rwn'], col['u_pct_gA'],
            col['u_ln']]
    missing = [c for c in need if c not in header]
    if missing:
        raise SystemExit(f"{cfg.icalc_file}[{cfg.icalc.sheet}]: no column(s) "
                         f"named {missing}")
    c_I, c_g, c_E, c_r, c_up, c_ul = (header[c] for c in need)
    cached_rows = cached.iter_rows(values_only=True)
    next(cached_rows)                       # the header
    old, new, other_formulas = [], [], set()
    for row, values in zip(range(2, ws.max_row + 1), cached_rows):
        gA, Eup, rwn = (ws.cell(row, c).value for c in (c_g, c_E, c_r))
        if gA is None or Eup is None or rwn is None:
            continue
        prev = ws.cell(row, c_I).value
        val = float(model_intensity(C, kT, float(gA), float(Eup), float(rwn)))
        ws.cell(row, c_I).value = val
        if isinstance(prev, (int, float)) and prev > 0:
            old.append(float(prev))
            new.append(val)
        # u_ln, and anything else still holding a formula, becomes a number
        u_pct = ws.cell(row, c_up).value
        if isinstance(ws.cell(row, c_ul).value, str) and u_pct is not None:
            ws.cell(row, c_ul).value = float(np.log1p(float(u_pct) / 100.0))
        for cell in ws[row]:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                other_formulas.add(cell.column_letter)
                cell.value = values[cell.column - 1]
    try:
        wb.save(cfg.icalc_file)
    except PermissionError:
        raise SystemExit(
            f"cannot write {cfg.icalc_file}: it is open in another program "
            f"(Excel locks the file). Close it and run again; the untouched "
            f"copy is at {backup}.")
    return backup, len(new), np.array(old), np.array(new), sorted(other_formulas)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument('--lines', default=None,
                    help='line_classifications.csv to fit (default: the '
                         'output_csv of the configuration)')
    ap.add_argument('--clip', type=float, default=2.0,
                    help='drop points with |ln(Icalc/Iobs)| above this before '
                         'the second fit (default 2)')
    ap.add_argument('--write', action='store_true',
                    help='recompute the Icalc column of the workbook')
    ap.add_argument('--plot', metavar='PNG', default=None,
                    help='save the Boltzmann plot to this file')
    ap.add_argument('--points', metavar='CSV', default=None,
                    help='save the fitted points to this file')
    a = ap.parse_args(argv)

    cfg = config.load()
    lines_csv = a.lines or cfg.output_csv
    df, n_all = load_points(cfg, lines_csv)
    print(f"{os.path.basename(lines_csv)}: {n_all} accepted identifications, "
          f"{len(df)} of them with a gA in "
          f"{os.path.basename(cfg.icalc_file)}")
    if len(df) < 10:
        raise SystemExit('too few points to fit')

    C1, kT1, a1, b1, rms1 = straight_line(df['Eup'].to_numpy(),
                                          df['y'].to_numpy())
    print(f"\nfirst fit  ({len(df)} points):  C = {C1:.6g}   "
          f"kT = {kT1:.6g} cm^-1 ({kT1 / 8065.544:.4g} eV)   "
          f"rms of ln I = {rms1:.4g}")

    df['Icalc_fit'] = model_intensity(C1, kT1, df['gA'], df['Eup'], df['rwn'])
    df['ln_ratio'] = np.log(df['Icalc_fit'] / df['obs_intens'])
    keep = df['ln_ratio'].abs() <= a.clip
    print(f"clip |ln(Icalc/Iobs)| > {a.clip:g}: "
          f"{int((~keep).sum())} points dropped, {int(keep.sum())} kept")

    k = df[keep]
    C2, kT2, a2, b2, rms2 = straight_line(k['Eup'].to_numpy(),
                                          k['y'].to_numpy())
    print(f"\nsecond fit ({len(k)} points):  C = {C2:.6g}   "
          f"kT = {kT2:.6g} cm^-1 ({kT2 / 8065.544:.4g} eV)   "
          f"rms of ln I = {rms2:.4g}")
    print(f"\nadopted model:  Icalc = {C2:.6g} * gA * (rwn/1e8) "
          f"* exp(-Eup/{kT2:.6g})")
    print(f"in the configuration now: C = {cfg.intensity_model['C']:g}, "
          f"kT = {cfg.intensity_model['kT']:g} cm^-1")

    df['Icalc_final'] = model_intensity(C2, kT2, df['gA'], df['Eup'],
                                        df['rwn'])
    r = np.log(df['Icalc_final'] / df['obs_intens'])
    print(f"\nover all {len(df)} points, ln(Icalc/Iobs): "
          f"mean {r.mean():+.4g}, rms {np.sqrt((r ** 2).mean()):.4g}, "
          f"|.| <= 2 for {int((r.abs() <= 2).sum())} "
          f"({(r.abs() <= 2).mean():.1%})")

    if a.points:
        cols = ['wn_obs', 'low_id', 'upp_id', 'obs_intens', 'rwn', 'gA', 'Eup',
                'y', 'Icalc_fit', 'Icalc_final', 'ln_ratio', 'new', 'grade']
        out = df[[c for c in cols if c in df.columns]].copy()
        out['used_in_second_fit'] = keep.astype(int)
        out.to_csv(a.points, index=False)
        print(f"\npoints written to {a.points}")

    if a.plot:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots(figsize=(7.5, 5.5))
        ax.plot(df.loc[~keep, 'Eup'], df.loc[~keep, 'y'], '.', ms=3,
                color='0.75', label=f'dropped (|ln ratio| > {a.clip:g})')
        ax.plot(k['Eup'], k['y'], '.', ms=3, color='tab:blue', label='fitted')
        xs = np.array([df['Eup'].min(), df['Eup'].max()])
        ax.plot(xs, a2 + b2 * xs, '-', color='tab:red',
                label=f'ln C = {a2:.3f}, kT = {kT2:.0f} cm$^{{-1}}$')
        ax.set_xlabel(r'$E_{\rm up}$  (cm$^{-1}$)')
        ax.set_ylabel(r'$\ln\,[\,I_{\rm obs}\,(10^8/{\rm rwn})/gA\,]$')
        ax.set_title('Boltzmann plot of the accepted identifications')
        ax.legend(loc='upper right', fontsize=9)
        fig.tight_layout()
        fig.savefig(a.plot, dpi=150)
        print(f"plot written to {a.plot}")

    if a.write:
        backup, n, old, new, formula_cols = rewrite_icalc(cfg, C2, kT2)
        ratio = new / old
        if formula_cols:
            print(f"\ncolumns {', '.join(formula_cols)} held formulas and now "
                  f"hold the numbers those formulas gave")
        print(f"\n{os.path.basename(cfg.icalc_file)}: {n} Icalc values "
              f"recomputed (previous file copied to "
              f"{os.path.basename(backup)})")
        print(f"  new/old ratio: min {ratio.min():.4g}, median "
              f"{np.median(ratio):.4g}, max {ratio.max():.4g}")
        print("  NOTE: [intensity_model] in the configuration must carry the "
              "C and kT above, or every later")
        print("        run will warn that the file and the configuration "
              "disagree.")
    return 0


if __name__ == '__main__':
    sys.exit(main())
